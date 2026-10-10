from pathlib import Path
import sys,json
import pandas as pd,numpy as np
R=Path(__file__).resolve().parents[1];O=R/'resultados/cierre_socher_20261009'
sys.path.insert(0,str(R/'scripts'));from comun import canon
x=pd.read_parquet(R/'datos/precios_hist/202108.parquet');x['barra']=x.barra.map(canon)
q=x[x.barra=='QUILLOTA______220'][['dia','hora','P']].rename(columns={'P':'E'})
z=x.merge(q,on=['dia','hora'],validate='many_to_one');z['D']=z.P-z.E;z['L']=z.E*(z.fp-1);z['C']=z.P-z.E*z.fp
bars=pd.read_csv(O/'barras_2020_2021.csv');keys=bars[bars.anio==2021][['barra','distrito_nombre']]
z=z.merge(keys,on='barra',validate='many_to_one');z=z[z.distrito_nombre.isin(pd.read_csv(O/'cambios_controles_2020_2021.csv').unidad.tolist()+['LAMPA-4'])]
u=z.groupby(['distrito_nombre','dia','hora'])[['P','D','L','C']].mean().reset_index();u['grupo']=np.where(u.distrito_nombre=='LAMPA-4','Lampa','Controles')
g=u.groupby(['grupo','dia','hora'])[['P','D','L','C']].mean().reset_index();g.to_csv(O/'agosto_2021_grupos_horarios.csv',index=False)
cs=g[g.grupo=='Controles'].sort_values('C',ascending=False);print(cs.head(12).to_string(index=False));day=int(cs.iloc[0].dia)
# Re-read original workbook prices independently, all 24 hours for max-C day.
from extraer_precios_historicos import open_book,selected
from io import BytesIO
import openpyxl
sel=next(r for r in selected() if int(r['mes'])==202108);name,data=open_book(sel['archivo_proyecto']);wb=openpyxl.load_workbook(BytesIO(data),read_only=True,data_only=True,keep_links=False);print('SHEETS',wb.sheetnames)
ws=wb['CMG Barras'];ws.reset_dimensions();targets=['LAMPA_________220','QUILLOTA______220']+keys[keys.distrito_nombre=='LA REINA-6'].barra.tolist();rec=[]
for n,values in enumerate(ws.iter_rows(values_only=True),1):
 if n==1:continue
 if values[1]==day and canon(values[3]) in targets:
  rec.append({'barra':canon(values[3]),'dia':day,'hora':int(values[2]),'P_original':values[4]})
 if len(rec)==24*len(targets):break
wb.close();raw=pd.DataFrame(rec).merge(x[['barra','dia','hora','P']],on=['barra','dia','hora'],validate='one_to_one');assert len(raw)==24*len(targets);err=float(abs(raw.P-raw.P_original).max());assert err<1e-9
raw.to_csv(O/'muestra_original_agosto_2021.csv',index=False)
res=dict(mes=202108,dia=day,barras=targets,celdas=len(raw),error_maximo=err,libro=name,horas_controles_C_mayor_1=int((cs.C>1).sum()),horas_totales=len(cs),C_promedio=float(cs.C.mean()),nota='C residual según factores por tramos; no identifica causa operativa ni toda congestión. Verificación directa de precios, no reextracción completa del libro.')
(O/'verificacion_agosto_2021.json').write_text(json.dumps(res,indent=2,ensure_ascii=False));print(res)
