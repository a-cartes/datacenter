"""Documento de cuatro páginas; resultados leídos de la ejecución de resultados_socher.R."""
from pathlib import Path
import csv
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "resultados/socher_20261005"
OUT = ROOT / "entregables/socher_2026"
OUT.mkdir(parents=True, exist_ok=True)
res = json.loads((DATA / "resultados.json").read_text())
rows = {r["especificacion"]: r for r in res["variantes"]}
main = rows["principal_nodos_incorporados"]
loo = {r["excluida"]: r for r in res["exclusion"]}
prospectiva = json.loads((DATA / "comparacion_2030.json").read_text())
escenarios = prospectiva['escenarios_ilustrativos']

def n(x, decimals=2):
    return f"{float(x):.{decimals}f}".replace(".", ",")

doc = Document()
doc.settings.odd_and_even_pages_header_footer = False
sec = doc.sections[0]
sec.different_first_page_header_footer = False
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.8)
sec.left_margin, sec.right_margin = Cm(2.1), Cm(2.1)
sec.header_distance, sec.footer_distance = Cm(.8), Cm(.8)
for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"]:
    st = doc.styles[name]
    st.font.name = "Times New Roman"
    st.font.color.rgb = RGBColor(0, 0, 0)
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    fonts = st._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "csTheme"]:
        fonts.attrib.pop(qn("w:" + attr), None)
for el in list(doc.styles.element.iter(qn("w:pBdr"))):
    el.getparent().remove(el)
normal = doc.styles["Normal"]
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.04
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.widow_control = True
doc.styles["Title"].font.size = Pt(16)
doc.styles["Title"].font.bold = True
doc.styles["Title"].paragraph_format.space_after = Pt(8)
doc.styles["Heading 1"].font.size = Pt(12)
doc.styles["Heading 1"].font.bold = True
doc.styles["Heading 1"].paragraph_format.space_before = Pt(9)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(5)
doc.styles["Heading 1"].paragraph_format.keep_with_next = True
doc.styles["Caption"].font.size = Pt(9.5)
doc.styles["Caption"].font.italic = False
doc.styles["Caption"].paragraph_format.space_after = Pt(5)
doc.styles["Caption"].paragraph_format.keep_with_next = True
doc.core_properties.title = "Centros de datos y diferencias territoriales de precios eléctricos en Santiago"
doc.core_properties.author = "Alan Cartes Fica; Ricardo Troncoso"
doc.core_properties.subject = "Resumen extendido para SOCHER 2026"
doc.core_properties.comments = "Resumen extendido para SOCHER 2026. Correo de correspondencia: alancartes@outlook.com."
header = sec.header.paragraphs[0]
header.text = "SOCHER 2026  |  Resumen extendido de investigación en curso"
header.style = doc.styles["Normal"]
for run in header.runs:
    run.font.size = Pt(8.5)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run("5 de octubre de 2026  |  ")
run.font.size = Pt(8.5)
field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

def para(text, bold=False, small=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text); r.bold = bold
    if small: r.font.size = Pt(9.5)
    return p

def heading(text):
    return doc.add_paragraph(text, style="Heading 1")

def table(headers, data, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, w in zip(t.columns, widths): c.width = Cm(w)
    for c, h, w in zip(t.rows[0].cells, headers, widths): c.text = h; c.width = Cm(w)
    for values in data:
        cells = t.add_row().cells
        for c, val, w in zip(cells, values, widths): c.text = str(val); c.width = Cm(w)
    for i, row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit"); pr.append(no_split)
        if i == 0:
            repeat = OxmlElement("w:tblHeader"); pr.append(repeat)
        for j, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for e in ["top", "left", "bottom", "right"]:
                x = OxmlElement("w:" + e)
                for k,v in {"val":"single","sz":"4","color":"D9D9D9"}.items(): x.set(qn("w:"+k),v)
                borders.append(x)
            cp.append(borders)
            margins = OxmlElement("w:tcMar")
            for e, val in [("top","60"),("bottom","60"),("left","75"),("right","75")]:
                x = OxmlElement("w:"+e); x.set(qn("w:w"),val); x.set(qn("w:type"),"dxa"); margins.append(x)
            cp.append(margins)
            if i == 0:
                sh = OxmlElement("w:shd"); sh.set(qn("w:fill"),"E8EDF1"); cp.append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.02
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs: run.font.size = Pt(9.5); run.bold = i==0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

# Página 1: problema, pregunta y diseño.
doc.add_paragraph("Centros de datos y diferencias territoriales de precios eléctricos en Santiago", style="Title")
p = doc.add_paragraph("Réplica empírica y comparación con Virginia para la planificación hacia 2030")
p.paragraph_format.space_after = Pt(7)
for r in p.runs: r.font.size = Pt(12)
para("Alan Cartes Fica y Ricardo Troncoso", bold=True)
para("Universidad del Bío-Bío, Chile. Correspondencia: alancartes@outlook.com.", small=True)
para("Palabras clave: centros de datos, precios eléctricos, transmisión, economía regional, diferencias en diferencias.", small=True)
heading("1 Introducción y justificación")
para('Chile ha convertido la expansión de los centros de datos en una apuesta de desarrollo económico. El Plan Nacional de Data Centers 2024-2030 (PDATA) busca consolidar al país como un polo tecnológico de América Latina y vincular la inversión con capacidades de inteligencia artificial (IA), investigación y servicios públicos (MinCiencia, 2024). La magnitud prevista exige anticipar sus efectos territoriales: las proyecciones divulgadas por Systep (2025) sitúan la demanda eléctrica nacional del sector en 325 MW en 2025 y 1.207 MW en 2030, con 1.038 MW concentrados en Santiago.')
para('Northern Virginia, el mayor mercado mundial de centros de datos, ofrece una referencia para examinar esa expansión (JLARC, 2024). La comparación permite estudiar cómo la concentración digital interactúa con la capacidad de abastecimiento local. Para Santiago, la pregunta económica es qué condiciones de red permiten capturar los beneficios de la inversión y cómo se distribuyen territorialmente sus costos eléctricos.')
para('La literatura identifica mecanismos que justifican esta oportunidad. La digitalización reduce costos de búsqueda, reproducción y verificación de información (Goldfarb y Tucker, 2019). Brynjolfsson, Li y Raymond (2025) documentan aumentos de productividad por asistencia de IA generativa, heterogéneos entre trabajadores. Estos resultados motivan el desarrollo de capacidades digitales, aunque los beneficios locales de alojar infraestructura dependen del acceso al cómputo, las competencias y los vínculos productivos.')
para('En Chile existen aplicaciones concretas en sectores estratégicos: Codelco (2026) utiliza analítica e IA para anticipar fallas y apoyar decisiones operacionales; la astronomía demanda procesamiento intensivo, como ilustra la infraestructura distribuida del Observatorio Rubin (2025). El PDATA también contempla una multi-nube estatal y capacidad de IA para investigación. Esta agenda tiene una dimensión estratégica: ampliar el acceso nacional al cómputo y fortalecer la gestión de datos públicos requiere gobernanza, ciberseguridad y capacidades institucionales, además de instalaciones físicas.')
para('El potencial renovable refuerza esta agenda. Chile encabezó en 2022 el ranking Climatescope de mercados emergentes para inversión eléctrica limpia, por segunda vez en tres años (BloombergNEF, 2022). Aprovechar recursos nacionales para digitalización y electrificación del transporte exige coordinar generación, almacenamiento y transmisión. Además, la relación entre servicios digitales y electricidad depende de la eficiencia tecnológica (Masanet et al., 2020). Por ello, el recurso renovable disponible debe analizarse junto con los costos de abastecer cada localización y horario.')
para('La expansión también plantea desafíos regulatorios. En julio de 2026, el SEA determinó que “Habilitación Data Center Alto Jahuel”, de DC Terranova SpA en Buin, no debía ingresar al SEIA según los antecedentes presentados, manteniendo las obligaciones sectoriales aplicables (SEA, 2026). Este caso refuerza la necesidad de información territorial que articule proyectos, conexiones y planificación eléctrica.')
para('En este contexto, adaptamos a Santiago el diseño de primeras entradas de Mamkhezri, Sun y Yang (2026) para Virginia. El aporte combina un inventario trazable con precios por barra y evidencia de restricciones de conexión. La réplica identifica una variación positiva, sensible a los comparadores, y permite delimitar qué evidencia histórica puede sustentar una evaluación prospectiva hacia 2030.')
heading("2 Objetivo e hipótesis")
para("El objetivo es estimar cómo cambia la diferencia local de precios eléctricos tras la primera entrada de un centro de datos, evaluar la comparabilidad con Virginia y dimensionar la exposición económica de la demanda proyectada a 2030. La hipótesis es que una carga concentrada puede ampliar esas diferencias cuando encuentra restricciones de abastecimiento local.")
heading("3 Datos y metodología")
para("Se integran precios horarios oficiales del Coordinador Eléctrico Nacional para 2017-2025, información técnica de barras y subestaciones, distritos y covariables del Censo 2017 y un inventario de 63 registros de instalaciones o etapas, respaldado por 138 hitos y 65 fuentes. El inventario distingue construcción, apertura y expansión. Las conexiones se consideran acreditadas únicamente cuando existe documentación; la proximidad geográfica por sí sola no identifica la barra de retiro.")
para("La unidad de estimación es distrito-año. El resultado D es el promedio territorial de la diferencia horaria entre el costo marginal de cada barra y Quillota 220 kV, expresado en USD/MWh. Se admiten barras con al menos 90% de cobertura anual, incluyendo las incorporadas durante el período. D mide diferencias espaciales relativas; su correspondencia con la suma de congestión y pérdidas de PJM es conceptual y no una identidad física.")

# Página 2: estimación y resultados con comprobaciones esenciales.
para("Se identifican tres primeras entradas: Colina-3 en 2021, Lampa-4 en 2022 y San Bernardo-11 en 2023. El año de Lampa se actualiza con la inauguración documentada por InvestChile (2026); no se interpreta como fecha exacta del primer retiro. De 58 controles candidatos, el emparejamiento probit con diez vecinos y reemplazo selecciona 18 controles únicos. Se emplean población, viviendas y composición demográfica y habitacional del Censo 2017; son adaptaciones declaradas de las covariables originales.")
para("El estimador de Sun y Abraham (2021) incluye efectos fijos de distrito y año y utiliza el año anterior a la entrada como referencia. Las copias de controles reciben peso inverso a su multiplicidad. La muestra representa 189 observaciones distrito-año únicas; las 297 filas con copias no añaden unidades independientes. Se agrupan errores por distrito y se calculan intervalos t consistentes con esa especificación. La actualización de fechas y la incorporación de barras son correcciones posteriores al protocolo inicial y se conservan las estimaciones anteriores.")
heading("4 Resultados preliminares")
para(f"La estimación principal es {n(main['att'])} USD/MWh, con intervalo convencional de 95% entre {n(main['ic95_inf'])} y {n(main['ic95_sup'])}. La reconstrucción por diferencias de cohorte y año reproduce el mismo coeficiente. Este resultado describe una variación positiva de precios relativos en la muestra, aunque su interpretación causal requiere considerar las diferencias previas y la exposición compartida de los controles.")
doc.add_paragraph("Tabla 1  Estimaciones y sensibilidades en USD por MWh", style="Caption")
spec_labels = [
 ("Principal con barras incorporadas", "principal_nodos_incorporados"),
 ("Solo barras presentes los nueve años", "sensibilidad_barras_fijas"),
 ("Tratamiento desde el año calendario siguiente", "sensibilidad_calendario_siguiente"),
 ("Controles con correlación previa menor a 0,5", "sensibilidad_controles_baja_correlacion"),
 ("Controles alejados más de 5 km", "sensibilidad_controles_lejos"),
]
table(["Especificación", "Estimación", "IC 95% convencional", "Tratados / controles"],
      [[label,n(rows[k]['att']),f"[{n(rows[k]['ic95_inf'])}; {n(rows[k]['ic95_sup'])}]",f"{rows[k]['tratadas']} / {rows[k]['controles']}"] for label,k in spec_labels],
      [7.8,2.0,4.2,2.8])
para("Nota: las alternativas cambian fechas, composición de barras o comparadores; no estiman necesariamente el mismo efecto. Los filtros de controles usan barras fijas y fechas corregidas. La última variante conserva solo un distrito tratado tras el emparejamiento. Los intervalos son convencionales y su cobertura no está garantizada con tan pocos tratados.", small=True)
para(f"Al excluir Lampa-4, la estimación cae a {n(loo['LAMPA-4']['att'])} USD/MWh, con intervalo [{n(loo['LAMPA-4']['ic95_inf'])}; {n(loo['LAMPA-4']['ic95_sup'])}]. Al excluir Colina-3 o San Bernardo-11 se obtienen, respectivamente, {n(loo['COLINA-3']['att'])} y {n(loo['SAN BERNARDO-11']['att'])}. Por tanto, Lampa concentra buena parte del resultado agregado. El filtro de baja correlación invierte el signo, lo que impide presentar el coeficiente principal como un efecto causal robusto.")

# Página 3: dinámica, comparación y alcance de la extensión.
doc.add_paragraph("Figura 1  Estudio de eventos de la especificación principal", style="Caption")
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
p.add_run().add_picture(str(DATA / "figura_eventos.png"), width=Cm(16.7))
para("Fuente: elaboración propia. Barras: intervalos t convencionales al 95%. Referencia: t = -1. La composición cambia con el horizonte: t = -6 y t = 4 incluyen una cohorte; t = -5 y t = 3 incluyen dos; los demás horizontes incluyen tres.", small=True)
para("Los coeficientes anteriores a la entrada no se mantienen alrededor de cero. En t = -2 la diferencia respecto del año de referencia alcanza 2,05 USD/MWh, una magnitud relevante frente al resultado posterior. Aunque las cohortes disponibles cambian entre horizontes, la evidencia obliga a revisar tendencias paralelas y posibles cambios de exposición eléctrica. La significancia convencional del promedio posterior no resuelve esta limitación.")
heading("5 Comparación con Virginia y extensión hacia 2030")
para("El estudio de Virginia reporta 2,49 USD/MWh para congestión más pérdidas, con 19 sectores tratados y 115 controles únicos. Santiago registra 1,05 USD/MWh sobre D y tres distritos tratados. La comparación es descriptiva y utiliza el resultado publicado, sin reestimar Virginia. La EIA (2026) registra allí casi 30 TWh adicionales de ventas comerciales entre 2019 y 2025, impulsadas principalmente por centros de datos; el agregado incluye otras actividades y 2025 es preliminar.")
para("Systep (2025), atribuyendo las cifras al Coordinador, reporta 325 MW nacionales para 2025 y 1.207 MW para 2030: un aumento de 271,4%. La RM concentraría 1.038 MW, el 86,0% del total. La tabla convierte esa potencia en energía bajo supuestos explícitos y valora su sensibilidad a un cambio uniforme de precio.")
doc.add_paragraph("Tabla 2  Energía y sensibilidad de valoración spot en la RM en 2030", style="Caption")
table(["Potencia media como fracción de 1.038 MW", "Energía anual TWh", "Millones USD por cada +1 USD/MWh"],
      [[n(r['fraccion_potencia_publicada']*100,0)+"%",n(r['energia_TWh']),n(r['millones_USD_por_1_USD_MWh'])] for r in escenarios],
      [6.8,4.0,6.0])
para("Cálculo: energía = 1.038 MW × fracción × 8.760 horas; valoración = energía × variación de precio. Las fracciones son ilustrativas. La nota fuente no precisa si los MW son potencia media o máxima; si ya son media, corresponde 100%. No son pronósticos de precios ni de gasto contractual.", small=True)

# Página 4: evaluación económica, conclusión y referencias.
para("La carta DE 04560-25 del Coordinador (2025) documenta el mecanismo territorial: Ascenty SCL4 proyecta 50 MW a fines de 2027 y 150 MW a fines de 2028. El informe aprobado propone Punta Verde 110 kV porque Chacabuco carece de capacidad suficiente bajo contingencia N-1. Este antecedente identifica demanda y restricción de red, sin acreditar operación futura ni un alza de precios ya ocurrida. El proyecto no se suma al agregado regional, evitando duplicidades.")
para("La evaluación prospectiva incorpora tres escenarios oficiales de planificación del Coordinador (2026), con 349 barras y 496 líneas por escenario. Se extrajeron demanda y costos marginales proyectados para 2030 en 120 bloques temporales anuales y se verificó la correspondencia de la demanda entre entradas y salidas oficiales. Las bases incluyen generación, hidrología, embalses, almacenamiento y transmisión. Proporcionan los insumos para contrastar una implementación del despacho y comparar escenarios con y sin expansión digital, manteniendo iguales las demás condiciones. La desagregación de las cargas digitales y la simulación contrafactual siguen pendientes.").paragraph_format.keep_together = True
para("La sensibilidad monetaria de la tabla expresa exposición bruta a precios mayoristas; los contratos pueden modificar el gasto efectivo. Los precios oficiales extraídos corresponden a escenarios base y no identifican el efecto de la expansión digital. El coeficiente histórico de primera entrada tampoco identifica una respuesta por MW y no se extrapola a 2030.", small=True)
heading("6 Conclusiones")
para("La réplica adaptada identifica una variación positiva de precios relativos después de las entradas, concentrada principalmente en Lampa y sensible a la selección de controles. Las diferencias previas y la escasez de distritos tratados impiden establecer por ahora una atribución causal firme. Estos hallazgos delimitan qué puede aprenderse del diseño de primera entrada en un mercado con instalaciones y expansiones concentradas territorialmente.")
para("El aporte para la ciencia regional combina una base territorial reproducible, evidencia de restricciones de conexión y una cuantificación condicional de exposición económica. Virginia ofrece una referencia para estudiar mecanismos; la evolución de Santiago dependerá de sus propias inversiones y condiciones de red. La evaluación ambiental constituye una extensión posterior.")
heading("Bibliografía básica")
references = [
"Coordinador Eléctrico Nacional. (2025, 25 de julio). Comunica obra urgente Nueva S/E Seccionadora Punta Verde 110 kV y aplicación de artículo 102°. Carta DE 04560-25. https://cartas.coordinador.cl/download_saved_file/6887c8cf3563571c24dc6f4a",
"Coordinador Eléctrico Nacional. (2026). Propuesta Final de Expansión de la Transmisión 2026. Anexo VII, bases y resultados PLP y PowerFactory [bases de datos]. https://www.coordinador.cl/desarrollo/documentos/propuesta-2026/propuesta-final-2026/anexos-propuesta-final-2026/",
"Coordinador Eléctrico Nacional. (s. f.). Costos marginales reales [base de datos, registros 2017-2025]. https://www.coordinador.cl/costos-marginales/",
"Energy Information Administration. (2026, 5 de mayo). Commercial electricity sales have soared in Virginia, driven by data centers. Today in Energy. https://www.eia.gov/todayinenergy/detail.php?id=67664",
"Instituto Nacional de Estadísticas. (2017). Censo de Población y Vivienda 2017 [datos y cartografía]. https://www.ine.gob.cl/herramientas/portal-de-mapas/geodatos-abiertos",
"InvestChile. (2026). Servicios globales e industria tecnológica en Chile (marzo, p. 14). Agencia de Promoción de la Inversión Extranjera. https://tools.investchile.gob.cl/hubfs/01-%20Descargables%20InvestChile/Ebooks%20sectoriales/Servicios%20Globales/InvestChile_ebook_servicios%20globales.pdf?hsLang=en-gb",
"Mamkhezri, J., Sun, X., y Yang, Y. (2026). The Hidden Cost of the Cloud: Data Centers and Transmission-Induced Electricity Prices (versión del 9 de marzo) [documento de trabajo]. SSRN. https://doi.org/10.2139/ssrn.5736562",
"Sun, L., y Abraham, S. (2021). Estimating dynamic treatment effects in event studies with heterogeneous treatment effects. Journal of Econometrics, 225(2), 175-199. https://doi.org/10.1016/j.jeconom.2020.09.006",
"Systep. (2025, 23 de septiembre). Data centers cuadruplicarán su demanda de energía a 2030: alza presionará al sistema eléctrico [nota que reproduce proyecciones del Coordinador]. https://systep.cl/data-centers-cuadruplicaran-su-demanda-de-energia-a-2030-alza-presionara-al-sistema-electrico/",
]
references.extend(['BloombergNEF. (2022). Climatescope 2022: Power transition factbook (pp. 47-48). https://www.global-climatescope.org/downloads/climatescope-2022-power-report-en.pdf', 'Brynjolfsson, E., Li, D., y Raymond, L. (2025). Generative AI at work. The Quarterly Journal of Economics, 140(2), 889-942. https://doi.org/10.1093/qje/qjae044', 'Codelco. (2026, 4 de agosto). División El Teniente desarrolla centro de datos para anticipar riesgos y optimizar procesos. https://www.codelco.com/division-el-teniente-desarrolla-centro-de-datos-para-anticipar-riesgos-y', 'Goldfarb, A., y Tucker, C. (2019). Digital economics. Journal of Economic Literature, 57(1), 3-43. https://doi.org/10.1257/jel.20171452', 'Joint Legislative Audit and Review Commission (JLARC). (2024). Data centers in Virginia. https://jlarc.virginia.gov/landing-2024-data-centers-in-virginia.asp', 'Masanet, E., Shehabi, A., Lei, N., Smith, S., y Koomey, J. (2020). Recalibrating global data center energy-use estimates. Science, 367(6481), 984-986. https://doi.org/10.1126/science.aba3758', 'Ministerio de Ciencia, Tecnología, Conocimiento e Innovación (MinCiencia). (2024). Plan Nacional de Data Centers 2024-2030. https://minciencia.gob.cl/uploads/filer_public/95/6b/956b8c9f-d937-4b4d-8f6c-a871495a52ff/plan_nacional_de_data_centers_pdata.pdf', 'Observatorio Rubin. (2025, 15 de mayo). Preparados, listos… ¡a procesar!: El Observatorio Rubin traerá una avalancha de datos cósmicos. https://rubinobservatory.org/es/news/rubin-data-deluge', 'Servicio de Evaluación Ambiental (SEA). (2026, 24 de julio). Pronunciamiento sobre Habilitación Data Center Alto Jahuel. PERTI-2026-5962. https://firma.sea.gob.cl/publicaciones/2026/07/24/1784914495_42EC06A7-D626-4DCD-8456-E599CAEE7C3F'])
references.sort(key=str.casefold)
for r in references:
    text, url = r.split("https://", 1)
    p = para(text, small=True)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"),p.part.relate_to("https://"+url,RT.HYPERLINK,is_external=True))
    rr = OxmlElement("w:r"); prop=OxmlElement("w:rPr")
    sz=OxmlElement("w:sz");sz.set(qn("w:val"),"19");prop.append(sz)
    col=OxmlElement("w:color");col.set(qn("w:val"),"000000");prop.append(col)
    rr.append(prop); tt=OxmlElement("w:t");tt.text="[Enlace]";rr.append(tt);link.append(rr);p._p.append(link)
    p.paragraph_format.left_indent = Cm(.45)
    p.paragraph_format.first_line_indent = Cm(-.45)
    p.paragraph_format.space_after = Pt(3)

path = OUT / "Resumen_extendido_SOCHER_2026.docx"
for el in list(doc.element.iter(qn("w:pBdr"))):
    el.getparent().remove(el)
doc.save(path)
# Texto de control y estado guardados fuera de la carpeta de entregables.
(DATA / "texto_resumen.txt").write_text('\n'.join(p.text for p in doc.paragraphs))
(DATA / "estado_documento.json").write_text(json.dumps({
    "fecha":"2026-10-05", "autores":["Alan Cartes Fica","Ricardo Troncoso"],
    "filiacion":"Universidad del Bío-Bío", "correo_correspondencia":"alancartes@outlook.com",
    "formato":"A4, Times New Roman 11, 4 páginas previstas; rango 2-4 aportado por usuario",
    "estado":"Versión de envío con resultados preliminares e identificación completa; no enviado",
    "faltantes_identificacion":[],
    "gestion_envio":"Recepción SOCHER fuera del plazo publicado no confirmada",
    "trabajo_en_curso":"Estimación contrafactual de precios y costos 2030 pendiente, declarada en el manuscrito",
    "datos":"Bases de Replica_RM_desde_cero; no se modificaron bases previas",
    "reproduccion":["Rscript scripts/resultados_socher.R", "python scripts/calculos_comparacion_2030.py", "python scripts/crear_resumen_socher.py"],
    "prospectiva":"Bases PET26 incorporadas: tres escenarios, 349 barras y 496 líneas; entradas y salidas de demanda 2030 conciliadas. Equivalencias de energía y sensibilidad de valoración spot; todavía no son efectos estimados de precios 2030",
},ensure_ascii=False,indent=2))
print(path)
