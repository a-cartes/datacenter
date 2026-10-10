"""Inspección de fuentes originales descargadas, sin ejecutar macros ni estimar efectos."""
from collections import Counter
import csv
import hashlib
from io import BytesIO, TextIOWrapper
import json
from pathlib import Path
import warnings
from zipfile import ZipFile

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "fuentes" / "precios"
SOURCES = {
    "Antecedentes_CMG_Real_def_240715.zip": {
        "url": "https://www.coordinador.cl/wp-content/uploads/2024/07/Antecedentes_CMG_Real_def_240715.zip",
        "periodo": "2024-07-15",
        "publicacion": "2024-07-24",
    },
    "cmg1712_def.zip": {
        "url": "https://www.coordinador.cl/wp-content/old-docs/2018/04/cmg1712_def.zip",
        "periodo": "2017-12",
        "publicacion": "2018-04-10",
    },
}


def workbook(zipped, name):
    return openpyxl.load_workbook(
        BytesIO(zipped.read(name)), read_only=True, data_only=True, keep_links=False
    )


def inspect_csv(zipped, name):
    bars, times, versions = Counter(), set(), Counter()
    keys = set()
    duplicates = missing = nonnumeric = count = 0
    price_field = "CMg[mills/kWh]" if name.startswith("Promedios") else "CMgR[mills/kWh]"
    with zipped.open(name) as stream:
        reader = csv.DictReader(TextIOWrapper(stream, encoding="utf-8-sig"), delimiter=";")
        fields = reader.fieldnames
        for row in reader:
            count += 1
            bars[row["BARRA"]] += 1
            key = (row["BARRA"], row["FECHA"], row["HORA"], row["MINUTO"])
            duplicates += key in keys
            keys.add(key)
            times.add(key[1:])
            value = row[price_field].strip()
            if not value:
                missing += 1
            else:
                try:
                    float(value)
                except ValueError:
                    nonnumeric += 1
            versions[row.get("REPORTE", "columnas E/P/R")] += 1
    return {
        "campos": fields, "filas": count, "barras_unicas": len(bars),
        "periodos_unicos": len(times), "filas_por_barra": dict(Counter(bars.values())),
        "claves_duplicadas": duplicates, "precios_vacios": missing,
        "precios_no_numericos": nonnumeric, "versiones": dict(versions),
    }


def inspect_recent(zipped):
    result = {name: inspect_csv(zipped, name) for name in zipped.namelist() if name.endswith(".csv")}
    for name in ["FPen_Mapeo_Barras.xlsx", "Fpen_Reales_240715.xlsx", "Barras_Subsistemas_20240715.xlsx"]:
        wb = workbook(zipped, name)
        sheets = {}
        for sheet in wb.worksheets:
            declared = [sheet.max_row, sheet.max_column]
            sheet.reset_dimensions()
            rows = sheet.iter_rows(values_only=True)
            header = next(rows)
            n = sum(1 for row in rows if any(v is not None for v in row))
            sheets[sheet.title] = {"cabecera": header, "filas_no_vacias_sin_cabecera": n, "dimension_declarada": declared}
        result[name] = sheets
        wb.close()
    return result


def inspect_historical(zipped):
    wb = workbook(zipped, "cmg1712_def.xlsm")
    result = {"hojas": wb.sheetnames}
    sheet = wb["CMG Barras dolares"]
    header = next(sheet.iter_rows(min_row=9, max_row=9, values_only=True))
    positions = [i for i, value in enumerate(header) if i >= 3 and isinstance(value, str) and value.strip()]
    bars = [header[i] for i in positions]
    times = set()
    missing = nonnumeric = 0
    for row in sheet.iter_rows(min_row=11, values_only=True):
        if row[0] != 201712:
            continue
        times.add(tuple(row[:3]))
        for i in positions:
            value = row[i]
            missing += value is None
            nonnumeric += value is not None and not isinstance(value, (int, float))
    result["CMG Barras dolares"] = {
        "fila_cabecera": 9, "columnas_barras": len(bars), "barras_unicas": len(set(bars)),
        "periodos_unicos": len(times), "celdas_precio_esperadas": len(positions) * len(times),
        "precios_vacios": missing, "precios_no_numericos": nonnumeric,
        "unidad_rotulada": "mills/kWh", "fecha_codificada": "Mes YYYYMM, día, hora",
    }
    result["Factores de Penalización"] = {
        "fila_cabecera": 10, "bloques_observados_en_cabecera": [1, 2, 3],
        "nota": "La correspondencia entre bloques y horas no se reconstruyó en esta inspección."
    }
    for name in ["central marginal por barra (a)", "central marginal por barra (b)", "Costo_Centrales"]:
        sheet = wb[name]
        width = 7 if name.startswith("central") else 9
        result[name] = {"campos_principales": next(sheet.iter_rows(max_row=1, max_col=width, values_only=True))}
    wb.close()
    return result


def main():
    report = {"fecha_revision": "2026-10-01", "metodo": "Inspección de archivos originales; no se ejecutan macros ni se recalculan libros.", "fuentes": {}}
    for name, source in SOURCES.items():
        path = DATA / name
        checksum = hashlib.file_digest(path.open("rb"), "sha256").hexdigest()
        with ZipFile(path) as zipped:
            entry = dict(source, bytes=path.stat().st_size, sha256=checksum)
            entry["miembros"] = [{"nombre": info.filename, "bytes": info.file_size} for info in zipped.infolist()]
            entry["inspeccion"] = inspect_recent(zipped) if name.startswith("Antecedentes") else inspect_historical(zipped)
            report["fuentes"][name] = entry
    output = DATA / "inspeccion_muestras.json"
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str) + "\n")
    print(output)
    for name, value in report["fuentes"].items():
        print(name, value["bytes"], value["sha256"])
        for member, detail in value["inspeccion"].items():
            if isinstance(detail, dict) and ("filas" in detail or "periodos_unicos" in detail):
                print(member, {k: v for k, v in detail.items() if k != "campos"})


if __name__ == "__main__":
    warnings.filterwarnings("ignore", message="Workbook contains no default style")
    main()
