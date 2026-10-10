"""Documento de cuatro páginas; resultados leídos de la ejecución de resultados_socher.R."""
from pathlib import Path
import csv
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "resultados/socher_20261005"
OUT = ROOT / "entregables/socher_2026"
OUT.mkdir(parents=True, exist_ok=True)
res = json.loads((DATA / "resultados.json").read_text())
rows = {r["especificacion"]: r for r in res["variantes"]}
main = rows["principal_nodos_incorporados"]
loo = {r["excluida"]: r for r in res["exclusion"]}
prospectiva = json.loads((DATA / "comparacion_2030.json").read_text())
escenarios = prospectiva['escenarios_ilustrativos']

def n(x, decimals=2):
    return f"{float(x):.{decimals}f}".replace(".", ",")

doc = Document()
doc.settings.odd_and_even_pages_header_footer = False
sec = doc.sections[0]
sec.different_first_page_header_footer = False
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.8)
sec.left_margin, sec.right_margin = Cm(2.1), Cm(2.1)
sec.header_distance, sec.footer_distance = Cm(.8), Cm(.8)
for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"]:
    st = doc.styles[name]
    st.font.name = "Times New Roman"
    st.font.color.rgb = RGBColor(0, 0, 0)
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    fonts = st._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "csTheme"]:
        fonts.attrib.pop(qn("w:" + attr), None)
for el in list(doc.styles.element.iter(qn("w:pBdr"))):
    el.getparent().remove(el)
normal = doc.styles["Normal"]
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.04
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.widow_control = True
doc.styles["Title"].font.size = Pt(16)
doc.styles["Title"].font.bold = True
doc.styles["Title"].paragraph_format.space_after = Pt(8)
doc.styles["Heading 1"].font.size = Pt(12)
doc.styles["Heading 1"].font.bold = True
doc.styles["Heading 1"].paragraph_format.space_before = Pt(9)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(5)
doc.styles["Heading 1"].paragraph_format.keep_with_next = True
doc.styles["Caption"].font.size = Pt(9.5)
doc.styles["Caption"].font.italic = False
doc.styles["Caption"].paragraph_format.space_after = Pt(5)
doc.styles["Caption"].paragraph_format.keep_with_next = True
doc.core_properties.title = "Centros de datos y diferencias territoriales de precios eléctricos en Santiago"
doc.core_properties.author = "Alan Cartes Fica; Ricardo Troncoso"
doc.core_properties.subject = "Resumen extendido para SOCHER 2026"
doc.core_properties.comments = "Resumen extendido para SOCHER 2026. Correo de correspondencia: alancartes@outlook.com."
header = sec.header.paragraphs[0]
header.text = "SOCHER 2026  |  Resumen extendido de investigación en curso"
header.style = doc.styles["Normal"]
for run in header.runs:
    run.font.size = Pt(8.5)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run("5 de octubre de 2026  |  ")
run.font.size = Pt(8.5)
field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

def para(text, bold=False, small=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text); r.bold = bold
    if small: r.font.size = Pt(9.5)
    return p

def heading(text):
    return doc.add_paragraph(text, style="Heading 1")

def table(headers, data, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, w in zip(t.columns, widths): c.width = Cm(w)
    for c, h, w in zip(t.rows[0].cells, headers, widths): c.text = h; c.width = Cm(w)
    for values in data:
        cells = t.add_row().cells
        for c, val, w in zip(cells, values, widths): c.text = str(val); c.width = Cm(w)
    for i, row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit"); pr.append(no_split)
        if i == 0:
            repeat = OxmlElement("w:tblHeader"); pr.append(repeat)
        for j, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for e in ["top", "left", "bottom", "right"]:
                x = OxmlElement("w:" + e)
                for k,v in {"val":"single","sz":"4","color":"D9D9D9"}.items(): x.set(qn("w:"+k),v)
                borders.append(x)
            cp.append(borders)
            margins = OxmlElement("w:tcMar")
            for e, val in [("top","60"),("bottom","60"),("left","75"),("right","75")]:
                x = OxmlElement("w:"+e); x.set(qn("w:w"),val); x.set(qn("w:type"),"dxa"); margins.append(x)
            cp.append(margins)
            if i == 0:
                sh = OxmlElement("w:shd"); sh.set(qn("w:fill"),"E8EDF1"); cp.append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.02
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs: run.font.size = Pt(9.5); run.bold = i==0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

# Cierre auditado 9 octubre. Conserva resultados y documentos anteriores.
import shutil
CLOSE=ROOT/'resultados/cierre_socher_20261009'
diag=json.loads((CLOSE/'diagnostico.json').read_text())
checks=json.loads((CLOSE/'contrastes_estadisticos.json').read_text())
old=checks['fecha_original_misma_muestra']
title='¿Dónde se conecta la inteligencia artificial?'
subtitle='Centros de datos, red eléctrica y territorio en la Región Metropolitana ante el horizonte 2030'
doc.core_properties.title=title
doc.core_properties.subject=subtitle
footer.paragraph_format.space_after=Pt(0)
footer.clear();footer.add_run('9 de octubre de 2026  |  ').font.size=Pt(8.5)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.add_paragraph(title,style='Title')
p=doc.add_paragraph(subtitle);p.paragraph_format.space_after=Pt(7)
for r in p.runs:r.font.size=Pt(12)
para('Alan Cartes Fica y Ricardo Troncoso',bold=True)
para('Universidad del Bío-Bío, Chile. Correspondencia: alancartes@outlook.com.',small=True)
para('Palabras clave: inteligencia artificial, centros de datos, costos marginales, transmisión eléctrica, planificación territorial.',small=True)
heading('1 Introducción y contribución')
para('La inteligencia artificial tiene una geografía física: sus servicios requieren centros de datos, conexiones y electricidad. Su expansión abre una pregunta para las regiones que aspiran a atraer infraestructura digital: ¿qué sabemos de la exposición eléctrica de esos territorios? La literatura identifica una tensión entre las mejoras de eficiencia y el crecimiento de la demanda computacional (Masanet et al., 2020; de Vries, 2023), cuyos efectos económicos dependen de la oferta y la infraestructura disponible (Bogmans et al., 2026).')
para('Chile busca consolidarse como polo tecnológico latinoamericano mediante el Plan Nacional de Data Centers 2024-2030 (MinCiencia, 2024). Según proyecciones del Coordinador recogidas por Systep (2025), la demanda sectorial nacional pasaría de 325 MW en 2025 a 1.207 MW en 2030, un incremento cercano al 271%. La Región Metropolitana (RM) concentraría 1.038 MW, aproximadamente el 86% del total. El desfase entre el desarrollo de estas instalaciones y las obras de transmisión convierte su localización en una decisión económica territorial. Estas cifras corresponden a potencia proyectada para el conjunto del sector.')
para('Este estudio establece una línea de base frente a esa expansión: reúne instalaciones, conexiones y precios por barra, y adapta a la RM un diseño aplicado en Virginia. Documenta diferencias entre inauguración y conexión eléctrica, y una estimación histórica de +1,05 USD/MWh sensible a Lampa y a la evolución previa de los controles. El aporte es mostrar qué revela la evidencia disponible sobre la exposición eléctrica de los territorios. Constituye la primera etapa de una agenda que evaluará escenarios hacia 2030 y, con mejores datos de operación, si existen efectos causales y bajo qué condiciones.')
heading('2 Datos y estrategia de análisis')
para('El diagnóstico integra un inventario trazable de 63 registros de instalaciones o etapas, 138 hitos y 65 fuentes, un panel de costos marginales de 2017-2025 y documentación de conexiones que incluye antecedentes de 2026. El inventario abarca distintos servicios digitales; la IA motiva el horizonte de expansión, mientras la unidad observada es la infraestructura de centros de datos. Se distinguen construcción, apertura, conexión y ampliación para reconstruir su incorporación territorial y eléctrica.')
para('La unidad del panel es el distrito censal de 2017 (INE, 2017). Las subestaciones se asignan a los distritos por coordenadas y sus barras mediante un diccionario técnico. El resultado D es la diferencia horaria entre el costo marginal local y Quillota 220 kV, agregada anualmente. Quillota sirve de referencia por su factor de penalización uno y su pertenencia al subsistema principal en las muestras oficiales verificadas. Se incorporan barras con cobertura anual mínima de 90%. D captura diferencias locales de precios que incluyen pérdidas y separación de subsistemas; su alcance difiere del componente de congestión y pérdidas de PJM.')
para('La réplica adaptada identifica primeras entradas en Colina-3 (2021), Lampa-4 (2022) y San Bernardo-11 (2023). La fecha de Lampa se revisó desde 2023 a la inauguración de 2022 documentada por InvestChile (2026); ambas fechas se comparan con la misma muestra. Un emparejamiento probit de diez vecinos con reemplazo, basado en población, viviendas y composición demográfica y habitacional, selecciona 18 controles únicos entre 58 candidatos. Se aplica Sun y Abraham (2021), con efectos fijos de distrito y año, referencia en el año previo y ponderación inversa a la multiplicidad de las copias. Las 297 filas representan 189 observaciones distrito-año únicas. Se examinan tendencias previas y sensibilidad a fechas y composición de la muestra (Roth et al., 2023); las revisiones se registran como posteriores a la estimación inicial.')
heading('3 Evidencia histórica de precios').paragraph_format.page_break_before = True
para(f'La estimación distrital principal alcanza {n(main["att"])} USD/MWh. Su sensibilidad es sustantiva: al excluir Lampa se reduce a {n(loo["LAMPA-4"]["att"])} USD/MWh y, al comparar cada cohorte con el promedio de todo su período previo, el contraste descriptivo es {n(diag["contraste_pre_completo"])} USD/MWh. La Tabla 1 conserva la fecha anterior de Lampa para distinguir el efecto de ese cambio del de otras decisiones de construcción del panel.')
doc.add_paragraph('Tabla 1  Contrastes históricos y sensibilidad de la estimación',style='Caption')
table(['Especificación','USD/MWh','IC 95% convencional'],[
 ['Principal distrital, Lampa 2022',n(main['att']),f'[{n(main["ic95_inf"])}; {n(main["ic95_sup"])}]'],
 ['Misma muestra, fecha anterior de Lampa 2023',n(old['ATT']),f'[{n(old["IC95"][0])}; {n(old["IC95"][1])}]'],
 ['Exclusión de Lampa',n(loo['LAMPA-4']['att']),f'[{n(loo["LAMPA-4"]["ic95_inf"])}; {n(loo["LAMPA-4"]["ic95_sup"])}]'],
 ['Contraste descriptivo contra todo el período previo',n(diag['contraste_pre_completo']),'No calculado']],[10,2.3,4.5])
para('Nota: errores agrupados por distrito; con tres tratados, la cobertura nominal de los intervalos convencionales no está garantizada. La última fila compara cada cohorte con su propio promedio previo, con los mismos controles; es un contraste descriptivo, no una estimación causal. Las cohortes de 2021, 2022 y 2023 pesan 5/12, 4/12 y 3/12, según sus años posteriores observados.',small=True)
para('El análisis de Lampa localiza la principal fragilidad de la comparación. Entre 2020 y 2021, D sube de 0,67 a 1,52 USD/MWh en Lampa y de 1,80 a 6,81 en sus controles: la brecha relativa cae 4,16 USD/MWh porque el incremento es mayor en estos últimos. Diecisiete de los dieciocho controles presentan aumentos. El bajo diferencial relativo del año de referencia refleja, por tanto, una trayectoria eléctrica distinta entre ambos grupos, aun cuando el precio absoluto de Lampa también aumenta.')
doc.add_paragraph('Figura 1  Lampa y sus controles en el período de observación',style='Caption')
p=doc.add_paragraph();pic=p.add_run().add_picture(str(CLOSE/'diagnostico_lampa.png'),width=Cm(16.7));p.paragraph_format.space_after=Pt(3)
pic._inline.docPr.set('descr','Dos paneles con los diferenciales anuales de Lampa y sus controles entre 2017 y 2025 y la brecha entre ambos; el salto de los controles en 2021 precede a la inauguración de 2022.')
para('Fuente: elaboración propia con precios del Coordinador. Mismos 18 controles del modelo principal; línea vertical entre 2021 y 2022. El panel derecho representa Lampa menos el promedio de controles.',small=True)
para('La auditoría confirma que se comparan las mismas dos barras de Lampa y 52 de los controles entre 2020 y 2021. En los controles, el incremento de D se descompone contablemente en 2,03 USD/MWh de pérdidas relativas y 2,98 de separación de subsistemas, utilizando factores de penalización por tramos. En agosto de 2021, D alcanza 21,12 USD/MWh en controles frente a 2,36 en Lampa. La comprobación independiente de 96 precios horarios de ese mes reproduce exactamente el libro original. Son variaciones observadas de precios cuya causa operativa específica queda por establecer.')
para('Los contrastes previos comunes a las tres cohortes son 1,23; 1,35 y 2,05 USD/MWh en los horizontes −4, −3 y −2 respecto del año omitido. El test conjunto de los doce términos previos por cohorte no es calculable por singularidad de la covarianza; el test informado (p < 0,001) usa tres promedios de horizontes comunes y está sujeto a la misma limitación de pocos tratados. En conjunto, los resultados restringen la atribución causal del +1,05 y muestran por qué el emparejamiento demográfico debe complementarse con comparabilidad eléctrica. La evidencia describe heterogeneidad territorial y tampoco establece un efecto nulo de la expansión digital.')
heading('4 Conexiones y diagnóstico territorial').paragraph_format.page_break_before = True
para('El inventario permite reconstruir una secuencia que una fecha única de apertura oculta. ODATA ST01 se inauguró en 2022, recibió en 2023 autorización para un proyecto de conexión de 30 MW en Nueva Lampa 220 kV y registra pruebas programadas para esas instalaciones en 2025 (InvestChile, 2026; Coordinador, 2023, 2025a). La previsión de demanda de la CNE, por su parte, identifica a ODATA Chile en Batuco 023 (CNE, 2026). Estos antecedentes obligan a distinguir etapas y puntos de abastecimiento antes de asignar retrospectivamente toda la potencia de un proyecto.')
para('La documentación de una etapa de EdgeConneX registra energización en mayo de 2024, mientras que, para las instalaciones eléctricas asociadas a ODATA ST02, el término de puesta en servicio informado corresponde a septiembre de 2026, posterior a la apertura comercial anunciada en 2023 (EdgeConneX Chile V, 2024; Lo Espejo Data Center, 2026). El hallazgo es de medición: apertura, conexión y utilización de capacidad representan dimensiones diferentes del desarrollo de una instalación. La potencia solicitada y la capacidad de transformación deben distinguirse de los retiros efectivamente medidos.')
para('El diagnóstico entrega una base para relacionar cada instalación con su exposición eléctrica y revela dos necesidades concretas: reconstruir conexiones vigentes por etapa y obtener series homogéneas de consumo. Los antecedentes disponibles aún no permiten estimar una respuesta por MW consumido en 2017-2025. Para la planificación territorial, esta distinción importa porque las fronteras administrativas y las áreas de influencia de la red no coinciden necesariamente.')
heading('5 Escala económica y horizonte 2030')
para('Virginia aporta un punto de comparación: Mamkhezri et al. (2026) estiman 2,49 USD/MWh en congestión y pérdidas con 19 territorios tratados y 115 controles únicos. Su experiencia permite formular mecanismos para Chile, considerando diferencias institucionales y de red. A su vez, Bogmans et al. (2026) muestran mediante simulaciones que la respuesta de oferta y transmisión condiciona el impacto de la demanda digital. Colangelo et al. (2026) demuestran flexibilidad de cargas de IA en una instalación experimental. Estas referencias orientan preguntas sobre localización, infraestructura y gestión de demanda.')
para('La restricción tiene una expresión local documentada: el Coordinador (2025b) plantea reforzar el abastecimiento de Ascenty SCL4 mediante Punta Verde 110 kV ante insuficiencia de Chacabuco bajo contingencia N-1, es decir, frente a la pérdida de un elemento de la red. El caso vincula expansión digital y requisitos de infraestructura en un territorio concreto; ilustra el valor de coordinar los plazos de inversión de ambos sectores.')
para('Como ilustración de escala, 1.038 MW equivalen a 5,46-9,09 TWh/año si la potencia media representa entre 60% y 100% de esa cifra. Cada USD/MWh adicional supondría 5,46-9,09 millones de USD anuales de valoración spot bruta. Este cálculo es una sensibilidad a una variación hipotética de precio, no un incremento estimado ni gasto contractual. La interpretación depende de la definición de potencia de la fuente; si corresponde a potencia media, aplica el extremo de 100%.')
heading('6 Conclusiones y agenda de investigación')
para('La competitividad digital de una región también se juega en su capacidad de abastecer electricidad donde y cuando se necesita. Este estudio aporta un diagnóstico trazable de la infraestructura de centros de datos de la RM, documenta diferencias entre apertura y conexión y muestra cómo las trayectorias eléctricas de los controles condicionan la lectura de los precios históricos. El aporte para la ciencia regional es vincular infraestructura digital, exposición eléctrica y decisiones de localización ante el horizonte 2030.')
para('La siguiente etapa evaluará escenarios de expansión mediante las bases oficiales de red y despacho PET26, considerando demanda y alternativas de abastecimiento; esa simulación está pendiente. La agenda doctoral contempla además obtener retiros medidos por instalación y seguir sus ampliaciones para investigar si existen efectos causales sobre los costos locales y bajo qué condiciones. Esta primera etapa establece la base empírica para anticipar decisiones de infraestructura y contrastar posteriormente sus consecuencias económicas.')
heading('Bibliografía básica').paragraph_format.page_break_before = True
refs=[
('Comisión Nacional de Energía. (2026). Anexos del Informe definitivo de previsión de demanda eléctrica 2025-2045, Anexo 2 [base de previsión].','https://www.cne.cl/wp-content/uploads/2026/02/Anexos-Informe-Definitivo-de-Prevision-de-Demanda-Electrica-2025-2045.xlsx'),
('Coordinador Eléctrico Nacional. (2023). Autorización de conexión Odata Datacenter ST01. Carta DE 02932-23, 3 de julio, pp. 1 y 4.','https://cartas.coordinador.cl/download_saved_file/64a328c53563574eab7a4eca'),
('Coordinador Eléctrico Nacional. (2025a). Estudio para análisis de falla EAF 312/2025. Anexo de solicitudes de trabajo, registro DI2025070947.','https://www.coordinador.cl/wp-content/uploads/2025/07/EAF-312-2025.pdf'),
('Coordinador Eléctrico Nacional. (2025b). Nueva S/E Seccionadora Punta Verde 110 kV. Carta DE 04560-25, 25 de julio.','https://cartas.coordinador.cl/download_saved_file/6887c8cf3563571c24dc6f4a'),
('Coordinador Eléctrico Nacional. (s. f.). Costos marginales reales [registros 2017-2025].','https://www.coordinador.cl/costos-marginales/'),
('EdgeConneX Chile V SpA. (2024). Solicitud de entrada en operación NUP 3975. Carta registrada DE 03409-24, 30 de mayo.','https://cartas.coordinador.cl/download_saved_file/6659e61a35635760544f0c17'),
('Instituto Nacional de Estadísticas. (2017). Censo de Población y Vivienda [datos y cartografía].','https://www.ine.gob.cl/herramientas/portal-de-mapas/geodatos-abiertos'),
('InvestChile. (2026). Servicios globales e industria tecnológica en Chile, marzo, p. 14.','https://tools.investchile.gob.cl/hubfs/01-%20Descargables%20InvestChile/Ebooks%20sectoriales/Servicios%20Globales/InvestChile_ebook_servicios%20globales.pdf?hsLang=en-gb'),
('Lo Espejo Data Center SpA. (2026). Solicitud de entrada en operación NUP 3129 Odata San Bernardo ST02. Carta del 25 de septiembre, registrada OP02215-26.','https://cartas.coordinador.cl/download_saved_file/6abad5b73563570a7f1eb54a'),
('Mamkhezri, J., Sun, X., y Yang, Y. (2026). The Hidden Cost of the Cloud: Data Centers and Transmission-Induced Electricity Prices. Documento de trabajo, versión marzo, 57 pp.','https://doi.org/10.2139/ssrn.5736562'),
('Ministerio de Ciencia, Tecnología, Conocimiento e Innovación. (2024). Plan Nacional de Data Centers 2024-2030.','https://minciencia.gob.cl/uploads/filer_public/95/6b/956b8c9f-d937-4b4d-8f6c-a871495a52ff/plan_nacional_de_data_centers_pdata.pdf'),
('Sun, L., y Abraham, S. (2021). Estimating dynamic treatment effects in event studies with heterogeneous treatment effects. Journal of Econometrics, 225(2), 175-199.','https://doi.org/10.1016/j.jeconom.2020.09.006'),
('Systep. (2025). Data centers cuadruplicarán su demanda de energía a 2030: alza presionará al sistema eléctrico, 23 de septiembre.','https://systep.cl/data-centers-cuadruplicaran-su-demanda-de-energia-a-2030-alza-presionara-al-sistema-electrico/')]
refs.extend([('Bogmans, C., Ganpurev, G., Gomez-Gonzalez, P., Melina, G., Pescatori, A., y Thube, S. (2026). Power hungry: How AI will drive energy demand. Energy Economics, 158, 109278.', 'https://doi.org/10.1016/j.eneco.2026.109278'), ('Colangelo, P., Coskun, A. K., Megrue, J., Roberts, C., Sengupta, S., Sivaram, V., Tiao, E., Vijaykar, A., Williams, C., Wilson, D. C., Records, B., MacFarland, Z., Dreiling, D., Morey, N., Ratnayake, A., y Vairamohan, B. (2026). AI data centres as grid-interactive assets. Nature Energy, 11, 254-261.', 'https://doi.org/10.1038/s41560-025-01927-1'), ('de Vries, A. (2023). The growing energy footprint of artificial intelligence [Comentario]. Joule, 7(10), 2191-2194.', 'https://doi.org/10.1016/j.joule.2023.09.004'), ('Masanet, E., Shehabi, A., Lei, N., Smith, S., y Koomey, J. (2020). Recalibrating global data center energy-use estimates [Policy Forum]. Science, 367(6481), 984-986.', 'https://doi.org/10.1126/science.aba3758'), ('Roth, J., Sant’Anna, P. H. C., Bilinski, A., y Poe, J. (2023). What’s trending in difference-in-differences? A synthesis of the recent econometrics literature. Journal of Econometrics, 235(2), 2218-2244.', 'https://doi.org/10.1016/j.jeconom.2023.03.008')])
refs.sort(key=lambda item: item[0].casefold())
alternativas={
 'https://doi.org/10.1016/j.eneco.2026.109278':('Versión de trabajo (FMI, 2025)', 'https://www.imf.org/en/Publications/WP/Issues/2025/04/21/Power-Hungry-How-AI-Will-Drive-Energy-Demand-566304'),
 'https://doi.org/10.1038/s41560-025-01927-1':('Preprint (2025)', 'https://arxiv.org/abs/2507.00909v1')}
def enlace_visible(p,url):
 link=OxmlElement('w:hyperlink')
 link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 rr=OxmlElement('w:r');pr=OxmlElement('w:rPr')
 sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');pr.append(sz)
 color=OxmlElement('w:color');color.set(qn('w:val'),'155A8A');pr.append(color)
 rr.append(pr);tt=OxmlElement('w:t');tt.text=url;rr.append(tt);link.append(rr);p._p.append(link)
for txt,url in refs:
 p=para(txt+' ',small=True);p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.space_after=Pt(3);p.paragraph_format.left_indent=Cm(.45);p.paragraph_format.first_line_indent=Cm(-.45)
 enlace_visible(p,url)
 if url in alternativas:
  label,alt=alternativas[url]
  p.add_run(' '+label+': ').font.size=Pt(9)
  enlace_visible(p,alt)
for el in list(doc.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
path=OUT/'Resumen_extendido_SOCHER_2026_revision_20261009.docx';doc.save(path)
(CLOSE/'texto_resumen.txt').write_text('\n'.join(p.text for p in doc.paragraphs))
(CLOSE/'estado_documento.json').write_text(json.dumps({'fecha':'2026-10-09','estado':'Cierre de resultados preliminares; no enviado','limite_paginas':'No verificado; cuatro páginas es objetivo editorial','fuentes_originales':'Conservadas','resultados_causales':'No atribución causal robusta','prospectiva':'Sensibilidades; no efectos de precios estimados2030'},ensure_ascii=False,indent=2))
print(path)
