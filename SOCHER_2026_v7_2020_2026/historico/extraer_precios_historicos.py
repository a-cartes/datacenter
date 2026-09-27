#!/usr/bin/env python3
"""Precios CMg anteriores al cambio de julio2024, desde originales CEN.

Extrae solo libros de precios y homologaciones de ZIP oficiales, incluso anidados.
Lee valores OOXML sin ejecutar macros. Mantiene etiquetas horarias originales,
identidad mensual y falta de homologación; no inventa fechas de tratamiento.
La selección de familias procede del inventario previo v6, no de los precios.
"""
from __future__ import annotations
import argparse, concurrent.futures, datetime as dt, hashlib, io, json, math
import re, shutil, sys, time, zipfile, xml.etree.ElementTree as ET
from pathlib import Path
import pandas as pd

HERE=Path(__file__).resolve().parent
BASE=HERE.parent
ROOT=BASE.parent
NS={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
GROUPS={
 'ALTO JAHUEL':['A.JAHUEL','A.JAHUEL2','AJAHUEL'], 'BAJOS DE MENA':['BJOSDMNA'],
 'CHENA':['C.CHENA','C.CHENA2'], 'CERRILLOS':['CERRILLOS'],
 'CHACABUCO':['CHACABUCO'], 'CHICUREO':['CHICUREO'], 'COLINA':['COLINA'],
 'CURAUMA':['CURAUMA'], 'EL SALTO':['EL_SALTO'], 'LAMPA':['LAMPA'],
 'NUEVA LAMPA':['NUEVA_LAMPA','NUEVALAMPA','N.LAMPA','NUELAMPA','NVALAMPA'],
 'LO CAMPINO':['LO_CAMPINO','LOCAMPINO'],
 'PUNTA VERDE':['PUNTA_VERDE','PUNTAVERDE','P.VERDE'],
 'LO AGUIRRE':['LO_AGUIRRE'], 'LO ESPEJO':['LO_ESPEJO'], 'MAIPO':['MAIPO'],
 'ODATA':['ODATA','ODATA_SAN_BERNARDO','ODATA_LAMPA'],
 'PADRE HURTADO':['P.HURTADO','PADRE_HURTADO'],
 'PANAMERICANA':['PANAMERICANA'], 'POLPAICO':['POLPAICO'],
 'QUILICURA':['QUILICURA'], 'RECOLETA':['RECOLETA'],
 'SAN BERNARDO':['SAN_BERNARDO'], 'SANTA MARTA':['SANTA_MARTA'],
 'SANTA TERESA':['SANTA_TERESA','SANTATERESA','STA.TERESA','S.TERESA'],
 'LA POLVORA':['LA_POLVORA'], 'CANTERAS':['CANTERAS'],
 'TAP LIBERTADORES':['TLIBERTADORES'], 'QUILLOTA':['QUILLOTA'],
 'CERRO NAVIA':['C.NAVIA'], 'LAS VEGAS':['L.VEGAS','L.VEGAS2'],
 'PUDAHUEL':['PUDAHUEL'], 'TAP QUILICURA':['TAP_QUILICURA'],
}
ALIASES={alias:group for group,items in GROUPS.items() for alias in items}

def sha(path):
 h=hashlib.sha256()
 with Path(path).open('rb') as f:
  for b in iter(lambda:f.read(2**20),b''):h.update(b)
 return h.hexdigest()

def log(d):print(json.dumps(d,ensure_ascii=False),flush=True)
def relative(p):return str(Path(p).relative_to(BASE))
def group_for(s):
 if not isinstance(s,str):return None
 return ALIASES.get(re.sub(r'_?\d{3}$','',s.strip()).rstrip('_'))
def col(s):
 n=0
 for c in s:
  if not c.isalpha():break
  n=n*26+ord(c)-64
 return n

def xml_rows(path,sheet_name):
 with zipfile.ZipFile(path) as z:
  strings=[]
  if 'xl/sharedStrings.xml' in z.namelist():
   with z.open('xl/sharedStrings.xml') as f:
    for _,e in ET.iterparse(f,events=('end',)):
     if e.tag.endswith('}si'):
      strings.append(''.join(t.text or '' for t in e.findall('.//m:t',NS)));e.clear()
  wb=ET.fromstring(z.read('xl/workbook.xml'))
  sheets={s.attrib['name']:s.attrib['{'+REL+'}id'] for s in wb.findall('.//m:sheet',NS)}
  rels=ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
  targets={r.attrib['Id']:r.attrib['Target'] for r in rels}
  target=targets[sheets[sheet_name]]
  target=target.lstrip('/') if target.startswith('/') else 'xl/'+target
  with z.open(target) as f:
   for _,e in ET.iterparse(f,events=('end',)):
    if e.tag.endswith('}row'):
     row={}
     for c in e:
      v=c.find('m:v',NS)
      if c.attrib.get('t')=='inlineStr':value=''.join(t.text or '' for t in c.findall('.//m:t',NS))
      elif v is None or v.text is None:continue
      elif c.attrib.get('t')=='s':value=strings[int(v.text)]
      else:value=v.text
      row[col(c.attrib['r'])]=value
     yield int(e.attrib['r']),row
     e.clear()

def save_member(z,info,out):
 data=z.read(info) # ZipFile verifica CRC antes de devolver el miembro.
 out.parent.mkdir(parents=True,exist_ok=True)
 if out.exists():
  if sha(out)!=hashlib.sha256(data).hexdigest():raise ValueError('Miembro divergente: '+str(out))
 else:out.write_bytes(data)
 return out

def unpack(source):
 match=re.search(r'cmg[_ ]?(\d{4})[_ -]?def',source.name,re.I)
 if not match:return []
 month='20'+match.group(1); found=[]; lineage=[]
 folder=HERE/'libros'/month
 def walk(z,parents,depth=0):
  if depth>3:raise ValueError('ZIP anidado a más de tres niveles')
  for info in z.infolist():
   name=Path(info.filename).name; low=name.lower()
   isbook=bool(re.match(r'cmg[_ ]?'+month[2:]+r'[_ -]?def\.(xlsm|xlsx)$',low))
   ishomo='homolog' in low and 'prg' not in low and low.endswith('.xlsx')
   if isbook or ishomo:
    p=save_member(z,info,folder/name);found.append(p)
    lineage.append({'ruta_zip':parents+[info.filename],'archivo':relative(p),'sha256':sha(p),'crc32':f'{info.CRC:08x}'})
   elif low.endswith('.zip') and 'cmg' in low and 'fpen' not in low:
    with zipfile.ZipFile(io.BytesIO(z.read(info))) as nested:walk(nested,parents+[info.filename],depth+1)
 with zipfile.ZipFile(source) as z:walk(z,[source.name])
 if found:
  (HERE/'manifiestos').mkdir(exist_ok=True)
  (HERE/'manifiestos'/f'extraccion_zip_{month}.json').write_text(json.dumps({'mes':month,'origen':str(source),'sha256_zip':sha(source),'miembros':lineage},ensure_ascii=False,indent=2))
 return found

def homologation(book,month):
 import openpyxl
 candidates=[p for p in book.parent.glob('*.xlsx') if 'homolog' in p.name.lower() and 'prg' not in p.name.lower() and (book.parent.name==month or month in p.name or month[2:] in p.name)]
 # Homologación v4 preservada con nombre uniforme; ninguna se aplica a otro mes.
 if not candidates:
  p=ROOT/'SOCHER_2026_v4/precios/originales'/f'homologacion_{month}_def.xlsx'
  if p.exists():candidates=[p]
 if len(candidates)>1:
  hashes={sha(p) for p in candidates}
  if len(hashes)!=1:raise ValueError('Homologación ambigua '+month)
 if not candidates:return {},None
 p=candidates[0];w=openpyxl.load_workbook(p,read_only=True,data_only=True)
 sheet=next((s for s in w.sheetnames if 'HOMOLOG' in s.upper()),w.sheetnames[0]);mapping={}
 for n,row in enumerate(w[sheet].values,1):
  if len(row)<3 or not group_for(row[0]):continue
  label=str(row[0]).strip()
  rec={'nombre_fisico':str(row[1] or ''),'codigo_fisico':str(row[2] or ''),'fila_homologacion':n}
  if label in mapping and mapping[label]!=rec:raise ValueError('Código mensual duplicado '+label)
  mapping[label]=rec
 w.close()
 return mapping,{'archivo':str(p),'sha256':sha(p),'hoja':sheet}

def process(book):
 start=time.monotonic(); month='20'+re.search(r'cmg[_ ]?(\d{4})',book.name,re.I).group(1)
 out=HERE/'derivados';out.mkdir(exist_ok=True)
 report=HERE/'manifiestos'/f'precios_{month}.json';report.parent.mkdir(exist_ok=True)
 checksum=sha(book);mapping,homo=homologation(book,month)
 codehash=sha(Path(__file__))
 if report.exists():
  prev=json.loads(report.read_text())
  if prev.get('libro_sha256')==checksum and prev.get('homologacion')==homo and prev.get('script_sha256')==codehash and (out/f'precios_{month}.parquet').exists():return {'mes':month,'estado':'ya_verificado','filas':prev['filas']}
 rows=[];blocks=[];missing=[];labels=set();seen=set();n_source=0
 for idx,row in xml_rows(book,'CMG Barras'):
  if not blocks:
   blocks=[c for c,v in row.items() if str(v).strip()=='Barra']
   if blocks:
    if not all('mills/kWh' in row.get(c+1,'') for c in blocks):raise ValueError('Unidad no reconocida '+str(book))
   continue
  for c in blocks:
   raw_label=row.get(c);family=group_for(raw_label)
   if not family:continue
   label=raw_label.strip();ym=int(float(row[c-3]));day=int(float(row[c-2]));hour=int(float(row[c-1]))
   if str(ym)!=month:raise ValueError('Mes ajeno al libro '+month)
   if not 1<=hour<=25:raise ValueError('Etiqueta horaria fuera1–25')
   date=dt.date(ym//100,ym%100,day).isoformat();n_source+=1
   if date>='2024-07-15':continue # Solo régimen histórico en esta tabla.
   key=(label,date,hour)
   if key in seen:raise ValueError('Observación repetida en fuente '+repr(key))
   seen.add(key);labels.add(label);raw=row.get(c+1)
   if raw is None:missing.append({'barra':label,'fecha':date,'hora':hour,'fila':idx});continue
   price=float(raw)
   if not math.isfinite(price):raise ValueError('Precio no finito')
   meta=mapping.get(label,{})
   rows.append({'fecha':date,'hora_etiqueta':hour,'barra_cmg':label,
    'barra_cmg_original':raw_label,'familia_barra':family,
    'codigo_fisico':meta.get('codigo_fisico',''),'nombre_fisico':meta.get('nombre_fisico',''),
    'precio_usd_mwh':price,'regimen_precio':'historico_antes_20240715',
    'estado_identidad':'homologacion_mensual_exacta' if meta.get('codigo_fisico') else 'etiqueta_sin_homologacion_mensual',
    'mes_fuente':month,'version_fuente':'DEF','archivo_fuente':str(book),
    'fila_fuente':idx,'columna_fuente':c,'semantica_hora':'etiqueta_fuente_1_25_sin_conversion_UTC'})
 if not rows:raise ValueError('Ningún precio pertinente '+str(book))
 df=pd.DataFrame(rows).sort_values(['fecha','hora_etiqueta','barra_cmg']).reset_index(drop=True)
 pq=out/f'precios_{month}.parquet';csv=out/f'precios_{month}.csv.gz'
 df.to_parquet(pq,index=False,compression='zstd');df.to_csv(csv,index=False,compression={'method':'gzip','mtime':0})
 cov=df.groupby(['barra_cmg','familia_barra','codigo_fisico','estado_identidad']).agg(filas=('precio_usd_mwh','size'),dias=('fecha','nunique'),desde=('fecha','min'),hasta=('fecha','max')).reset_index()
 cov.to_csv(out/f'cobertura_{month}.csv',index=False)
 result={'mes':month,'filas':len(df),'barras_etiqueta':len(labels),'fecha_min':df.fecha.min(),'fecha_max':df.fecha.max(),'dias':df.fecha.nunique(),
  'fuente_filas_seleccionadas_antes_corte':n_source,'libro':str(book),'libro_sha256':checksum,'homologacion':homo,
  'estado_identidad':df.estado_identidad.value_counts().to_dict(),'faltantes_explicitos':missing,
  'duplicados':0,'outputs':{relative(p):sha(p) for p in [pq,csv]},'script_sha256':codehash,
  'ejecucion_utc':dt.datetime.now(dt.timezone.utc).isoformat(),'segundos':round(time.monotonic()-start,1)}
 report.write_text(json.dumps(result,ensure_ascii=False,indent=2));return result

def discover(months=None):
 errors=[]
 for source in sorted((HERE/'originales').glob('*.zip')):
  match=re.search(r'cmg[_ ]?(\d{4})[_ -]?def',source.name,re.I)
  month='20'+match.group(1) if match else None
  if months and month not in months:continue
  try:unpack(source)
  except Exception as e:
   error={'archivo':str(source),'mes':month,'error_tipo':type(e).__name__,'error':str(e)}
   errors.append(error);log({'error_extraccion_zip':error})
 if errors:
  folder=HERE/'manifiestos';folder.mkdir(exist_ok=True)
  stamp=dt.datetime.now(dt.timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
  (folder/f'errores_extraccion_{stamp}.json').write_text(json.dumps(errors,ensure_ascii=False,indent=2))
 books={}
 for p in sorted((HERE/'libros').rglob('*.xlsm')):
  m=re.match(r'cmg[_ ]?(\d{4})[_ -]?def\.xlsm$',p.name,re.I)
  if m:books['20'+m.group(1)]=p
 for p in (ROOT/'SOCHER_2026_v4/precios/originales').glob('cmg*_def.xlsm'):
  month='20'+re.search(r'cmg(\d{4})',p.name).group(1)
  if month not in books:books[month]=p
 return books

def main():
 ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--months',nargs='*');ap.add_argument('--workers',type=int,default=3);ap.add_argument('--unpack',type=Path)
 a=ap.parse_args()
 if a.unpack:log({'extraidos':[str(p) for p in unpack(a.unpack)]});return
 books=discover(a.months);books={m:p for m,p in books.items() if m>='202001' and m<='202407' and (not a.months or m in a.months)}
 log({'libros_disponibles':sorted(books),'workers':a.workers})
 errors=[{'mes':m,'error':'No hay libro de precios disponible'} for m in (a.months or []) if m not in books]
 for error in errors:log(error)
 with concurrent.futures.ProcessPoolExecutor(max_workers=a.workers) as pool:
  pending={pool.submit(process,p):m for m,p in sorted(books.items())}
  for fut in concurrent.futures.as_completed(pending):
   try:log(fut.result())
   except Exception as e:
    err={'mes':pending[fut],'error':str(e)};errors.append(err);log(err)
 if errors:sys.exit(1)

if __name__=='__main__':main()
