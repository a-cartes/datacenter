"""Conserva antecedentes Dataluna 2020–2022 y extrae canales horarios de 2023.

Las unidades ambiguas y convenciones de signo distintas se mantienen separadas
de los CSV de 15 minutos. Ausencia de una etiqueta no constituye consumo cero.
"""
from pathlib import Path
from datetime import datetime, timedelta
import argparse
import hashlib
import json
import posixpath
import re
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
import pandas as pd

BASE = Path(__file__).resolve().parents[1]
ROOT = BASE.parents[1]
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def sha(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(8 * 1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def old_dataluna():
    source = ROOT / "SOCHER_2026_v3/consumo/dataluna_horario_2020_2022.csv.gz"
    manifest = ROOT / "SOCHER_2026_v3/consumo/manifest_fuentes_dataluna.csv"
    previous = pd.read_csv(manifest)
    checked = []
    for row in previous.itertuples():
        path = ROOT / row.archivo
        digest = sha(path)
        assert digest == row.sha256
        checked.append({"archivo": row.archivo, "sha256": digest, "bytes": path.stat().st_size})
    data = pd.read_csv(source, keep_default_na=False)
    assert len(data) == 20447 and not data.unidad_verificada.any()
    data["consumo_mw_homologado"] = float("nan")
    data["estado_enlace"] = "UNIDAD_SIGNO_Y_CORRESPONDENCIA_HISTORICA_PENDIENTES"
    data["archivo_extracto_previo"] = str(source.relative_to(ROOT))
    data["sha256_extracto_previo"] = sha(source)
    data.to_csv(BASE / "tablas/dataluna_horario_2020_2022_no_homologado.csv.gz", index=False, compression={"method": "gzip", "mtime": 0})
    data.to_parquet(BASE / "tablas/dataluna_horario_2020_2022_no_homologado.parquet", index=False, compression="zstd")
    return {"filas": len(data), "desde_local": data.fecha_local.min(), "hasta_local": data.fecha_local.max(),
            "originales_verificados": checked, "unidad_verificada": False,
            "nota": "Los 36 libros Norte originales cubren 2020–2022; DATALUNA aparece desde septiembre 2020. No se interpreta como entrada ni se imputa cero enero–agosto 2020."}


def first_sheet(z):
    book = ET.fromstring(z.read("xl/workbook.xml"))
    sheet = book.find("s:sheets/s:sheet", NS)
    relation = sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    target = next(r.attrib["Target"] for r in rels if r.attrib["Id"] == relation)
    member = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
    return member, sheet.attrib["name"]


def extract_book(path, sockets, legacy=False):
    w = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = w.worksheets[0]
    header = list(sheet.iter_rows(max_row=10, values_only=True))
    sheet_name = sheet.title
    w.close()
    assert str(header[2][5]).strip() in ("PUNTO MEDIDA", "CLAVE TF", "CLAVE_TF", "CLAVE DE TRANSFERENCIA")
    selected = []
    for i, socket in enumerate(header[2]):
        if socket in sockets:
            if not legacy:
                assert header[9][i] == "kWhD" and header[9][i + 1] == "kWhR", (path, socket)
            for j in (i,) if legacy else (i, i + 1):
                selected.append({"IDSOCKET": socket, "columna_excel": j + 1,
                    "columna_letras": openpyxl.utils.get_column_letter(j + 1), "canal_origen": header[9][j] or "valor_firmado_sin_unidad_rotulada",
                    "total_mensual_origen": header[8][j], "coordinado": header[3][i],
                    "subestacion": header[5][i], "tension_origen": str(header[6][i])})
    digest = sha(path)
    path_rel = str(path.relative_to(ROOT))
    lookup = {x["columna_letras"]: x for x in selected}
    target_cols = {"B", "C", "D", "E", "F", *lookup}
    state, records, skipped = {}, [], []
    with zipfile.ZipFile(path) as z:
        member, name = first_sheet(z)
        assert name == sheet_name
        strings = []
        if "xl/sharedStrings.xml" in z.namelist():
            shared = ET.fromstring(z.read("xl/sharedStrings.xml"))
            strings = ["".join(t.text or "" for t in si.iter(f"{{{NS['s']}}}t")) for si in shared]
        with z.open(member) as stream:
            for _, row in ET.iterparse(stream, events=("end",)):
                if row.tag != f"{{{NS['s']}}}row":
                    continue
                ri = int(row.attrib["r"])
                if ri <= 10:
                    row.clear()
                    continue
                cells = {}
                for c in row:
                    ref = c.attrib.get("r", "")
                    col = re.sub(r"[0-9]+", "", ref)
                    if col not in target_cols:
                        continue
                    value = c.find("s:v", NS)
                    if value is not None:
                        text = value.text
                        cells[col] = strings[int(text)] if c.attrib.get("t") == "s" else text
                for col, name in [("B", "anio"), ("C", "mes"), ("D", "dia")]:
                    if cells.get(col) is not None:
                        state[name] = cells[col]
                try:
                    y, month, day = (int(float(state[k])) for k in ("anio", "mes", "dia"))
                    date = datetime(y, month, day)
                    assert y == 2023 and month == int(path.parent.parent.name[:2])
                    uh = int(float(cells["F"]))
                    h = int(float(cells["E"])) if cells.get("E") is not None else None
                    assert 0 <= uh < 24 and (h is None or 0 <= h < 24)
                except (KeyError, ValueError, AssertionError):
                    skipped.append(ri)
                    row.clear()
                    continue
                local = date.replace(hour=h) if h is not None else None
                offset = (uh - h) % 24 if h is not None else None
                utc = local + timedelta(hours=offset) if offset in (3, 4) else None
                for col, info in lookup.items():
                    text = cells.get(col)
                    try:
                        value = float(text) if text is not None else float("nan")
                    except ValueError:
                        value = float("nan")
                    records.append({"archivo_original": path_rel, "sha256_original": digest,
                        "hoja": sheet_name, "fila_excel": ri, "IDSOCKET": info["IDSOCKET"],
                        "columna_excel": info["columna_excel"], "canal_origen": info["canal_origen"],
                        "unidad_rotulada": None if legacy else "kWh", "fecha_local": date.strftime("%Y-%m-%d"),
                        "hora_local_origen": h, "hora_utc_origen": uh,
                        "etiqueta_local": local.isoformat() if local else None,
                        "marca_utc_reconstruida": utc.isoformat() + "Z" if utc else None,
                        "utc_es_reconstruido": utc is not None, "valor_firmado_origen": value,
                        "texto_valor_original": text, "duracion_etiquetas_minutos": 60,
                        "estado_convencion_signo": "PRESERVADA_SIN_CONVERTIR_A_RETIRO",
                        "inicio_fin_fisico_acreditado": False, "enlace_precio_habilitado": False,
                        "fecha_tratamiento_validada": None})
                row.clear()
    data = pd.DataFrame(records)
    checks = []
    for info in selected:
        s = data.loc[data.columna_excel.eq(info["columna_excel"]), "valor_firmado_origen"]
        total = float(s.sum(min_count=1))
        expected = info["total_mensual_origen"]
        error = total - float(expected)
        assert abs(error) < 0.01, (path, info, error)
        checks.append({**info, "filas": len(s), "suma_reconstruida": total, "error_total_mensual": error})
    report = {"archivo_original": path_rel, "sha256": digest, "filas_extraidas": len(data),
              "canales_seleccionados": checks, "filas_fuera_calendario": skipped,
              "nota": "Cada socket tiene D/R separados; R puede estar firmado negativo en el libro. No se impone convención de CSV posterior."}
    print(path_rel + f": {len(data)} valores horario-canal, {len({x['IDSOCKET'] for x in selected})} identificadores", flush=True)
    return data, report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--solo-legacy-2023", action="store_true")
    args = parser.parse_args()
    if args.solo_legacy_2023:
        legacy_path = next((ROOT / "2023/01_Enero/Definitivo").glob("*Transferencia_Norte_*.xlsx"))
        legacy, legacy_check = extract_book(legacy_path, {"DATALUNA"}, legacy=True)
        assert len(legacy) == 744 and legacy.unidad_rotulada.isna().all()
        legacy.to_parquet(BASE / "tablas/dataluna_horario_2023_01_no_homologado.parquet", index=False, compression="zstd")
        legacy.to_csv(BASE / "tablas/dataluna_horario_2023_01_no_homologado.csv.gz", index=False, compression={"method": "gzip", "mtime": 0})
        report_path = BASE / "verificacion_antecedentes_horarios.json"
        report = json.loads(report_path.read_text())
        report["dataluna_2023_01"] = legacy_check
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        return
    (BASE / "tablas").mkdir(parents=True, exist_ok=True)
    report = {"dataluna_2020_2022": old_dataluna()}
    sockets = set(pd.read_csv(ROOT / "SOCHER_2026_v6_multicentros/mapeo/socket_candidatos_extracto_20260924.csv").punto_medida)
    parts, notes = [], []
    for p in sorted((ROOT / "2023").glob("*/Definitivo/*Puntos*.xlsx")):
        d, check = extract_book(p, sockets)
        parts.append(d)
        notes.append(check)
    d = pd.concat(parts, ignore_index=True)
    assert not d.duplicated(["archivo_original", "hoja", "fila_excel", "columna_excel"]).any()
    d.to_parquet(BASE / "tablas/mediciones_canales_horarios_2023_01_04.parquet", index=False, compression="zstd")
    d.to_csv(BASE / "tablas/mediciones_canales_horarios_2023_01_04.csv.gz", index=False, compression={"method": "gzip", "mtime": 0})
    report["libros_2023"] = notes
    report["filas_canales_2023"] = len(d)
    legacy_path = next((ROOT / "2023/01_Enero/Definitivo").glob("*Transferencia_Norte_*.xlsx"))
    legacy, legacy_check = extract_book(legacy_path, {"DATALUNA"}, legacy=True)
    assert len(legacy) == 744 and legacy.unidad_rotulada.isna().all()
    legacy.to_parquet(BASE / "tablas/dataluna_horario_2023_01_no_homologado.parquet", index=False, compression="zstd")
    legacy.to_csv(BASE / "tablas/dataluna_horario_2023_01_no_homologado.csv.gz", index=False, compression={"method": "gzip", "mtime": 0})
    report["dataluna_2023_01"] = legacy_check
    (BASE / "verificacion_antecedentes_horarios.json").write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(json.dumps({"dataluna_horas": report["dataluna_2020_2022"]["filas"], "canales_horarios_2023": len(d)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
