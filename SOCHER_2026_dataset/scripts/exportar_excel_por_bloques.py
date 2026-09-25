"""Exporta el panel y su diccionario a Excel mediante escritura secuencial.

Lee los bloques creados por preparar_exportacion_excel_csv.py, conserva los
tipos declarados y escribe todas las observaciones en el directorio indicado.
"""
import csv
import argparse
from datetime import datetime
import json
import os
from pathlib import Path
import warnings
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo, TableColumn
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--output-dir', default=os.getenv('SOCHER_OUTPUT_DIR', 'artifacts'),
                    help='Entrada preparada y Excel final. Ruta relativa a la raíz del repositorio; también SOCHER_OUTPUT_DIR.')
args = parser.parse_args()
requested = Path(args.output_dir).expanduser()
OUT = requested.resolve() if requested.is_absolute() else (ROOT / requested).resolve()
META = json.loads((OUT / 'preparacion/metadata.json').read_text())
WB = Workbook(write_only=True)
PANEL = WB.create_sheet('Panel horario')
DICT = WB.create_sheet('Diccionario')
FONT = Font(name='Arial', size=10, color='203040')
HEADER_FONT = Font(name='Arial', size=10, color='FFFFFF', bold=True)
HEADER_FILL = PatternFill('solid', fgColor='203B53')
HEADER_ALIGNMENT = Alignment(vertical='center', wrap_text=True)

def header(ws, names):
    cells = []
    for name in names:
        cell = WriteOnlyCell(ws, name)
        cell.font, cell.fill, cell.alignment = HEADER_FONT, HEADER_FILL, HEADER_ALIGNMENT
        cells.append(cell)
    ws.append(cells)

def add_table(ws, name, names, row_count):
    ref = f'A1:{get_column_letter(len(names))}{row_count+1}'
    table = Table(displayName=name, ref=ref)
    table.tableColumns = [TableColumn(id=i, name=h) for i, h in enumerate(names, 1)]
    table.autoFilter = AutoFilter(ref=ref)
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    with warnings.catch_warnings():
        warnings.filterwarnings('ignore', message='In write-only mode you must add table columns manually')
        ws.add_table(table)

PANEL.sheet_view.showGridLines = False
PANEL.sheet_properties.tabColor = '203B53'
PANEL.freeze_panes = 'E2'
PANEL.sheet_format.defaultRowHeight = 20
PANEL.row_dimensions[1].height = 56
for i, name in enumerate(META['headers']):
    width = 24
    if i == 0: width = 14
    elif i == 1: width = 16
    elif 'archivo' in name: width = 65
    elif 'nota' in name: width = 70
    elif 'nemotecnico' in name or 'clave_barra' in name: width = 43
    elif any(x in name for x in ('estado', 'regimen', 'vinculo')): width = 38
    elif 'id' in name: width = 32
    PANEL.column_dimensions[get_column_letter(i+1)].width = width
header(PANEL, META['headers'])
count = 0
for filename in META['chunks']:
    for row in json.loads((OUT / 'preparacion' / filename).read_text()):
        cells = []
        for name, value in zip(META['headers'], row):
            if value is None:
                cells.append(None)
                continue
            kind = META['types'][name].lower()
            if kind.startswith('datetime'):
                value = datetime.fromisoformat(value)
            cell = WriteOnlyCell(PANEL, value)
            cell.font = FONT
            if kind.startswith('datetime'): cell.number_format = 'yyyy-mm-dd'
            elif kind.startswith('float'): cell.number_format = '0.000000'
            elif kind.startswith('int'): cell.number_format = '0'
            elif isinstance(value, str): cell.number_format = '@'
            cells.append(cell)
        PANEL.append(cells)
        count += 1
    if count % 25000 == 0 or count == META['rows']:
        print(f'Filas escritas: {count}', flush=True)
assert count == META['rows']
add_table(PANEL, 'PanelHorario', META['headers'], count)

DICT.sheet_view.showGridLines = False
DICT.sheet_properties.tabColor = '448583'
DICT.freeze_panes = 'A2'
DICT.sheet_format.defaultRowHeight = 48
DICT.row_dimensions[1].height = 28
for col, width in [('A',42),('B',106),('C',18),('D',23),('E',15)]:
    DICT.column_dimensions[col].width = width
names = ['variable', 'definicion', 'unidad', 'tipo', 'n_nulos']
header(DICT, names)
for row in META['definitions']:
    cells = []
    for key in names:
        value = int(row[key]) if key == 'n_nulos' else (row[key] or None)
        cell = WriteOnlyCell(DICT, value)
        cell.font = Font(name='Arial', size=11, color='203040')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        if key == 'n_nulos': cell.number_format = '#,##0'
        cells.append(cell)
    DICT.append(cells)
add_table(DICT, 'DiccionarioPanel', names, len(META['definitions']))
WB.save(OUT / 'Dataset_articulo.xlsx')
print('Excel completo guardado.', flush=True)
