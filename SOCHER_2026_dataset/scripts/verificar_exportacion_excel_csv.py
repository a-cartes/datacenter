"""Verifica, sin escribir libros, las exportaciones completas CSV y Excel.

Compara cada celda con el CSV.gz fuente, usando su diccionario de tipos. Las
fechas nativas de Excel y las cadenas ISO son equivalentes. La tolerancia
numérica es absoluta (1e-10), no relativa. Requiere únicamente stdlib y
openpyxl; el libro se abre read_only=True y nunca se guarda desde este script.
"""
from pathlib import Path
import argparse
import csv
from collections import Counter
from datetime import date, datetime, timezone
import gzip
import hashlib
from itertools import zip_longest
import json
import math
import os
import re
import sys
import unicodedata
import zipfile
import xml.etree.ElementTree as ET

B = Path(__file__).resolve().parents[1]
ROOT = B.parent
DEFAULT_OUTPUT = os.getenv("SOCHER_OUTPUT_DIR", "artifacts")
SOURCE = B / "tablas/panel_barra_hora.csv.gz"
DICTIONARY = B / "documentacion/diccionario_panel.csv"
REPORT = ROOT / "artifacts/verificacion_exportacion_excel_csv.json"
TOLERANCE = 1e-10
EXPECTED_ROWS = 178016
EXPECTED_COLUMNS = 66


def sha(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def text_header(value):
    return "".join(c for c in unicodedata.normalize("NFKD", str(value or "").strip().lower())
                   if not unicodedata.combining(c))


def empty(value):
    return value is None or value == ""


def iso_date(value):
    if isinstance(value, datetime):
        if value.time().isoformat() == "00:00:00" and value.tzinfo is None:
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value)
    parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    if parsed.time().isoformat() == "00:00:00" and parsed.tzinfo is None:
        return parsed.date().isoformat()
    return parsed.isoformat()


def compare_value(expected, actual, dtype, excel):
    """Retorna (coincide, error numérico o None). No considera vacío == 0."""
    if empty(expected):
        return empty(actual), None
    if empty(actual):
        return False, None
    kind = dtype.lower()
    if kind.startswith("datetime"):
        try:
            return iso_date(expected) == iso_date(actual), None
        except (ValueError, TypeError):
            return False, None
    if kind in ("bool", "boolean"):
        if excel:
            return type(actual) is bool and actual == (expected.lower() == "true"), None
        return str(actual).lower() == expected.lower(), None
    if kind.startswith(("int", "uint", "float")):
        if isinstance(actual, bool):
            return False, None
        try:
            x, y = float(expected), float(actual)
            error = abs(x - y)
            return math.isfinite(x) and math.isfinite(y) and error <= TOLERANCE, error
        except (ValueError, TypeError, OverflowError):
            return False, None
    # Incluye identificadores, códigos, meses codificados y fila_homologacion,
    # originalmente de tipo texto: Excel debe preservar ese tipo también.
    return isinstance(actual, str) and expected == actual, None


class Audit:
    def __init__(self):
        self.errors = 0
        self.examples = []

    def error(self, message, **context):
        self.errors += 1
        if len(self.examples) < 25:
            self.examples.append(dict(mensaje=message, **context))

    def condition(self, valid, message, **context):
        if not valid:
            self.error(message, **context)


def compare_rows(audit, exported_rows, headers, types, label, excel=False):
    initial_errors = audit.errors
    missing = object()
    source_nulls = Counter()
    export_nulls = Counter()
    ids = set()
    export_count = source_count = paired_count = 0
    max_error = 0.0
    max_spread_error = {"q110": 0.0, "q220": 0.0}
    idx = {name: pos for pos, name in enumerate(headers)}
    with gzip.open(SOURCE, "rt", encoding="utf-8-sig", newline="") as f:
        source_rows = csv.reader(f)
        assert next(source_rows) == headers
        for line, (expected, actual) in enumerate(zip_longest(source_rows, exported_rows, fillvalue=missing), 2):
            if expected is not missing:
                source_count += 1
            if actual is not missing:
                export_count += 1
            if expected is missing or actual is missing:
                audit.error("Diferente cantidad de filas", archivo=label, fila=line)
                continue
            if len(actual) != len(headers):
                audit.error("Diferente cantidad de columnas", archivo=label, fila=line, n_columnas=len(actual))
                continue
            paired_count += 1
            oid = actual[idx["observacion_id"]]
            if oid in ids:
                audit.error("Identificador de observación duplicado", archivo=label, fila=line, observacion_id=oid)
            ids.add(oid)
            for col, want, got in zip(headers, expected, actual):
                source_nulls[col] += int(empty(want))
                export_nulls[col] += int(empty(got))
                valid, error = compare_value(want, got, types[col], excel)
                if error is not None and math.isfinite(error):
                    max_error = max(max_error, error)
                if not valid:
                    audit.error("Valor o tipo incompatible con la fuente", archivo=label, fila=line,
                                variable=col, esperado=want, obtenido=str(got), tipo_obtenido=type(got).__name__)
            for tag in ("q110", "q220"):
                price = actual[idx["precio_usd_mwh"]]
                reference = actual[idx[f"precio_usd_mwh_{tag}"]]
                spread = actual[idx[f"diferencial_{tag}_usd_mwh"]]
                if any(empty(x) for x in (price, reference, spread)):
                    continue
                try:
                    residual = abs(float(price) - float(reference) - float(spread))
                    max_spread_error[tag] = max(max_spread_error[tag], residual)
                    audit.condition(math.isfinite(residual) and residual <= TOLERANCE,
                                    "Diferencial distinto de precio menos referencia", archivo=label, fila=line, referencia=tag)
                except (ValueError, TypeError):
                    audit.error("Diferencial no numérico", archivo=label, fila=line, referencia=tag)
    audit.condition(source_count == EXPECTED_ROWS, "La fuente no tiene el tamaño esperado", filas=source_count)
    audit.condition(export_count == source_count, "Conteos diferentes", archivo=label,
                    filas_fuente=source_count, filas_exportadas=export_count)
    audit.condition(dict(source_nulls) == dict(export_nulls), "Los nulos cambiaron", archivo=label)
    return dict(estado="OK" if audit.errors == initial_errors else "FALLO",
                filas_fuente=source_count, filas_exportadas=export_count, filas_comparadas=paired_count,
                identificadores_unicos=len(ids), max_error_numerico_absoluto=max_error,
                max_residuo_diferencial=max_spread_error, nulos_fuente=dict(source_nulls), nulos_exportados=dict(export_nulls))


def check_dictionary(audit, sheet, source_dictionary, headers):
    rows = sheet.iter_rows(values_only=True)
    raw_headers = next(rows)
    positions = {text_header(v): i for i, v in enumerate(raw_headers)}
    audit.condition("variable" in positions and "definicion" in positions,
                    "Diccionario sin encabezados variable/definicion", hoja=sheet.title)
    if "variable" not in positions or "definicion" not in positions:
        return {"estado": "FALLO"}
    definitions = {}
    for line, row in enumerate(rows, 2):
        if not any(not empty(value) for value in row):
            continue
        variable = row[positions["variable"]]
        definition = row[positions["definicion"]]
        audit.condition(variable not in definitions, "Variable duplicada en diccionario Excel", variable=variable, fila=line)
        definitions[variable] = definition
    expected = {r["variable"]: r["definicion"] for r in source_dictionary}
    audit.condition(set(definitions) == set(headers), "Diccionario no cubre exactamente las 66 columnas")
    for variable, definition in expected.items():
        audit.condition(definitions.get(variable) == definition, "Definición diferente de la fuente", variable=variable)
    return {"filas": len(definitions), "variables_cubren_panel": set(definitions) == set(headers),
            "definiciones_identicas": definitions == expected}


def check_structure(audit, path):
    """Lee metadatos ZIP/XML; el XML grande se recorre por bloques, sin cargarlo."""
    initial_errors = audit.errors
    namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    expected_tables = {"PanelHorario": "A1:BN178017", "DiccionarioPanel": "A1:E67"}
    result = {"tablas": {}, "hojas": {}}
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if name.startswith("xl/tables/") and name.endswith(".xml"):
                root = ET.fromstring(z.read(name))
                filt = root.find("s:autoFilter", namespace)
                table = {"rango": root.get("ref"), "filtro": filt.get("ref") if filt is not None else None}
                result["tablas"][root.get("name")] = table
        audit.condition(set(result["tablas"]) == set(expected_tables), "Tablas Excel inesperadas")
        for name, ref in expected_tables.items():
            table = result["tablas"].get(name, {})
            audit.condition(table.get("rango") == ref and table.get("filtro") == ref,
                            "Rango o filtro de tabla incorrecto", tabla=name)
        for name, expected_pane in [("xl/worksheets/sheet1.xml", "E2"), ("xl/worksheets/sheet2.xml", "A2")]:
            pane = None
            table_parts = []
            declared_count = None
            # Etiquetas cortas de metadatos: conserva cola suficiente si una
            # etiqueta queda cortada entre bloques y evita duplicar coincidencias.
            tail = b""
            offset = 0
            seen = set()
            with z.open(name) as f:
                for chunk in iter(lambda: f.read(1024 * 1024), b""):
                    block = tail + chunk
                    base = offset - len(tail)
                    for match in re.finditer(rb"<(?:pane|tableParts|tablePart)\b[^>]*>", block):
                        position = base + match.start()
                        if position in seen:
                            continue
                        seen.add(position)
                        fragment = match.group().decode("utf-8")
                        if not fragment.endswith("/>"):
                            fragment = fragment[:-1] + "/>"
                        wrapper = '<root xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">' + fragment + '</root>'
                        element = ET.fromstring(wrapper)[0]
                        if element.tag == "pane":
                            pane = element.attrib
                        elif element.tag == "tableParts":
                            declared_count = element.get("count")
                        elif element.tag == "tablePart":
                            table_parts.append(element.attrib)
                    offset += len(chunk)
                    tail = block[-2048:]
            result["hojas"][name] = {"panel_congelado": pane, "tableParts_declaradas": declared_count,
                                       "tableParts_encontradas": len(table_parts)}
            audit.condition(pane is not None and pane.get("topLeftCell") == expected_pane and pane.get("state") == "frozen",
                            "Panel congelado inesperado", hoja=name)
            audit.condition(declared_count == "1" and len(table_parts) == 1,
                            "Hoja sin exactamente una tabla vinculada", hoja=name)
    result["estado"] = "OK" if initial_errors == audit.errors else "FALLO"
    return result


def main():
    global REPORT
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT,
                        help="Directorio de exportaciones e informe; relativo a la raíz del repositorio. También SOCHER_OUTPUT_DIR.")
    parser.add_argument("--xlsx", type=Path, help="Ruta explícita del Excel; reemplaza el archivo predeterminado del directorio de salida.")
    parser.add_argument("--csv", type=Path, help="Ruta explícita del CSV; reemplaza el archivo predeterminado del directorio de salida.")
    parser.add_argument("--report", type=Path, help="Ruta explícita del informe, opcional.")
    parser.add_argument("--panel-sheet", default="Panel horario")
    parser.add_argument("--dictionary-sheet", default="Diccionario")
    args = parser.parse_args()
    out = Path(args.output_dir).expanduser()
    out = out.resolve() if out.is_absolute() else (ROOT / out).resolve()
    args.xlsx = args.xlsx or out / "Dataset_articulo.xlsx"
    args.csv = args.csv or out / "Dataset_articulo.csv"
    REPORT = args.report or out / "verificacion_exportacion_excel_csv.json"
    audit = Audit()
    report = dict(fecha_utc=datetime.now(timezone.utc).isoformat(), tolerancia_absoluta=TOLERANCE,
                  modo="solo_lectura_de_exportaciones", script=str(Path(__file__).relative_to(ROOT)),
                  hash_script=sha(Path(__file__)), python=sys.version)
    paths = {"fuente_csv_gz": SOURCE, "diccionario_fuente": DICTIONARY,
             "exportacion_csv": args.csv, "exportacion_xlsx": args.xlsx}
    try:
        for label, path in paths.items():
            audit.condition(path.is_file(), "Archivo faltante", archivo=label, ruta=str(path))
        if audit.errors:
            return finish(audit, report)
        report["archivos"] = {k: {"ruta": str(p.relative_to(ROOT)) if p.is_relative_to(ROOT) else str(p),
                                  "sha256": sha(p), "bytes": p.stat().st_size} for k, p in paths.items()}
        report["estructura_excel"] = check_structure(audit, args.xlsx)
        with gzip.open(SOURCE, "rt", encoding="utf-8-sig", newline="") as f:
            headers = next(csv.reader(f))
        audit.condition(len(headers) == EXPECTED_COLUMNS and len(set(headers)) == EXPECTED_COLUMNS,
                        "Encabezado fuente inesperado", columnas=len(headers))
        with DICTIONARY.open(encoding="utf-8-sig", newline="") as f:
            dictionary = [r for r in csv.DictReader(f) if r["tabla"] == "panel_barra_hora"]
        types = {r["variable"]: r["tipo"] for r in dictionary}
        audit.condition(set(types) == set(headers), "Diccionario fuente incompatible con panel")
        if audit.errors:
            return finish(audit, report)
        with args.csv.open(encoding="utf-8-sig", newline="") as f:
            rows = csv.reader(f)
            header_csv = next(rows)
            audit.condition(header_csv == headers, "Encabezados CSV diferentes o reordenados")
            if header_csv == headers:
                report["csv"] = compare_rows(audit, rows, headers, types, "CSV")
        import openpyxl
        report["openpyxl"] = openpyxl.__version__
        workbook = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
        try:
            report["hojas"] = workbook.sheetnames
            audit.condition(args.panel_sheet in workbook.sheetnames, "Falta hoja del panel", hoja=args.panel_sheet)
            audit.condition(args.dictionary_sheet in workbook.sheetnames, "Falta hoja de diccionario", hoja=args.dictionary_sheet)
            if args.panel_sheet in workbook.sheetnames:
                sheet = workbook[args.panel_sheet]
                rows = sheet.iter_rows(values_only=True)
                header_xlsx = list(next(rows))
                audit.condition(header_xlsx == headers, "Encabezados Excel diferentes o reordenados")
                if header_xlsx == headers:
                    report["excel"] = compare_rows(audit, rows, headers, types, "Excel", excel=True)
            if args.dictionary_sheet in workbook.sheetnames:
                report["diccionario_excel"] = check_dictionary(audit, workbook[args.dictionary_sheet], dictionary, headers)
        finally:
            workbook.close()
    except Exception as exc:
        audit.error("Excepción en verificación", tipo=type(exc).__name__, detalle=str(exc))
    return finish(audit, report)


def finish(audit, report):
    report.update(estado="OK" if not audit.errors else "FALLO", n_errores=audit.errors,
                  ejemplos_errores=audit.examples, fecha_fin_utc=datetime.now(timezone.utc).isoformat())
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"estado": report["estado"], "errores": audit.errors,
                      "filas_csv": report.get("csv", {}).get("filas_exportadas"),
                      "filas_excel": report.get("excel", {}).get("filas_exportadas"),
                      "informe": str(REPORT)}, ensure_ascii=False))
    return 0 if not audit.errors else 1


if __name__ == "__main__":
    raise SystemExit(main())
