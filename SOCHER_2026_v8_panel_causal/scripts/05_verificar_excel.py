#!/usr/bin/env python3
"""Contrasta cada celda de datos del XLSX exportado contra su especificación.
openpyxl se usa exclusivamente para lectura, no para crear el libro.
"""
from pathlib import Path
import argparse
import datetime as dt
import hashlib
import json
import math
import openpyxl

p=argparse.ArgumentParser(); p.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[2]);p.add_argument('--xlsx',type=Path,required=True);a=p.parse_args()
stage=a.root.resolve()/'SOCHER_2026_v8_panel_causal'
spec=json.loads((stage/'revision/contenido_excel.json').read_text())
wb=openpyxl.load_workbook(a.xlsx,read_only=True,data_only=True)
assert wb.sheetnames==[s['name'] for s in spec['sheets']]
counts=[]
for definition in spec['sheets']:
    ws=wb[definition['name']];n=len(definition['columns'])
    assert list(next(ws.iter_rows(min_row=5,max_row=5,max_col=n,values_only=True)))==[c['label'] for c in definition['columns']]
    cellcount=0;maxerr=0.
    for index,(actual,expected) in enumerate(zip(ws.iter_rows(min_row=6,max_row=5+len(definition['rows']),max_col=n,values_only=True),definition['rows']),start=6):
        for j,c in enumerate(definition['columns']):
            ref=expected.get(c['key']);v=actual[j]
            if ref is None or ref=='': assert v in (None,''),(ws.title,index,c['key'],v,ref)
            elif c.get('type')=='date': assert isinstance(v,(dt.date,dt.datetime)) and v.isoformat()[:10]==ref,(ws.title,index,c['key'],v,ref)
            elif isinstance(ref,(int,float)) and not isinstance(ref,bool):
                assert isinstance(v,(int,float)) and math.isclose(v,ref,rel_tol=1e-12,abs_tol=1e-10),(ws.title,index,c['key'],v,ref)
                maxerr=max(maxerr,abs(v-ref))
            else: assert v==ref,(ws.title,index,c['key'],v,ref)
            cellcount+=1
    counts.append({'hoja':ws.title,'filas_datos':len(definition['rows']),'celdas_contrastadas':cellcount,'error_maximo_numerico':maxerr})
report={'estado':'OK','sha256_xlsx':hashlib.file_digest(a.xlsx.open('rb'),'sha256').hexdigest(),'hojas':counts,'alcance':'Todas las celdas del área de datos contra JSON; revisión visual por hoja en qa/. No es prueba en Excel de escritorio.'}
(stage/'validacion/verificacion_excel.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(report,ensure_ascii=False))
